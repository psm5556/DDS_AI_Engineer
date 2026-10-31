"""프로젝트 보고서(HTML·Word): 내용이 무엇이든 같은 틀(표지·목차·본문·맺음말, 머리글·꼬리말)에 담아 DDS Conversa가 발행한 문서임을 보인다.
AI가 쓴 보고서(models.Report)는 제목과 본문만 주고, 없으면 게시된 결과를 차례로 모은다. 틀은 코드가 붙이므로 내용으로 지울 수 없다.
그래프는 matplotlib으로 다시 그린다(웹 화면의 Plotly는 브라우저 전용)."""
from __future__ import annotations

import base64
import html
import io
import logging
import math
import re
import xml.etree.ElementTree as etree
from dataclasses import dataclass
from datetime import datetime
from functools import cache
from pathlib import Path

import matplotlib
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

from . import models, report_md  # noqa: E402
from .config import settings  # noqa: E402
from .report_md import TablePlan, TocItem, fence_states, lines_of, plain, plan_table, safe_text  # noqa: E402

log = logging.getLogger(__name__)
ISSUER = "DDS Conversa"
ASSETS = Path(__file__).parent / "assets"  # 서버에 frontend/public이 없을 수 있어 사본을 둔다(test_server_assets_equal_frontend_public)

# 문서 색은 앱의 밝은 테마 토큰(tokens.css)과 같다(test_report_colors_match_tokens)
COLOR = {"ink": "#1c1c1a", "ink-2": "#45443f", "ink-3": "#6a6860", "ink-4": "#9a988e", "line": "#d8d6ce", "line-2": "#e6e4dc",
         "panel": "#fbfaf6", "panel-2": "#f0eee8", "brand": "#c2380d", "logo-text": "#1428a0", "ok-ink": "#1f6a3d", "err-ink": "#9b0f28",
         "warn-ink": "#7a4f00", "print-bg": "#ffffff"}
# 그래프 계열 순서도 앱의 chartColors()와 같다: 주홍 → 청록 → 보라 → 호박 → 회청 → 빨강
PALETTE = ["#c2380d", "#0d9488", "#7c3aed", "#e0a100", "#64748b", "#b3122f"]
KIND_ORDER = {"markdown": 0, "table": 1, "scatter": 2, "line": 2, "bar": 2, "parallel": 2, "image": 3}
STATUS = {"complete": "완료", "failed": "실패", "in_progress": "진행 중", "pending": "대기", "skipped": "건너뜀"}
NO_IMAGE, NO_CHART = "(그림 파일을 찾을 수 없습니다)", "(이 결과는 그림으로 만들지 못했습니다)"
NO_WORD_IMAGE = "(이 그림 형식은 Word에 넣을 수 없습니다 → HTML로 받아 보세요)"


def _korean_font():
    # Linux 서버: fonts-nanum 또는 Noto CJK를 설치한다. Noto CJK 묶음 파일(ttc)은 이름이 JP로 읽혀도 한글이 들어 있다
    for name in ("Malgun Gothic", "NanumGothic", "NanumBarunGothic", "AppleGothic", "Noto Sans CJK KR", "Noto Sans KR",
                 "Noto Sans CJK JP"):
        if any(f.name == name for f in font_manager.fontManager.ttflist):
            plt.rcParams["font.family"] = name
            break
    plt.rcParams["axes.unicode_minus"] = False


_korean_font()


def _t(text):
    """그래프 글자의 $를 글자로 둔다. matplotlib은 $…$를 수식으로 읽어, 풀지 못하면 문서 전체가 멈춘다."""
    return None if text is None else str(text).replace("$", r"\$")


def _set_xticks(fig, ax, ticks: list[float], names: list) -> None:
    """가로축 이름을 겹치지 않게 둔다: 그대로 → 기울이기 → 세우기 → 건너뛰기. 어느 막대인지 못 읽는 그래프는 문서에 둘 수 없다."""
    names = [_t(n) for n in names]
    ax.set_xticks(ticks, names)
    if len(ticks) < 2:
        return
    renderer = fig.canvas.get_renderer()
    boxes = [t.get_window_extent(renderer) for t in ax.get_xticklabels()]
    gap = ax.get_window_extent(renderer).width / len(ticks)
    high = max(b.height for b in boxes)
    if max(b.width for b in boxes) * 1.1 <= gap:
        return
    names = [n if len(n) <= 18 else n[:17] + "…" for n in names]
    if gap >= high * 1.5:
        ax.set_xticks(ticks, names, rotation=45, ha="right", rotation_mode="anchor")
    else:
        step = math.ceil(high * 1.15 / gap)
        ax.set_xticks(ticks[::step], names[::step], rotation=90)


def _has_data(a: models.Artifact) -> bool:
    p = a.payload or {}
    if a.kind in ("scatter", "line"):
        return any(s.get("points") for s in p.get("series") or [])
    if a.kind == "bar":
        return bool(p.get("categories")) and any(s.get("values") for s in p.get("series") or [])
    return True


def chart_png(a: models.Artifact) -> bytes | None:
    p = a.payload or {}
    series = p.get("series") or []
    if not _has_data(a):  # 빈 축만 있는 그림에 번호를 붙이지 않는다
        return None
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=130)
    try:
        if a.kind in ("scatter", "line"):
            for i, s in enumerate(series):
                pts = s.get("points") or []
                xs, ys = [q["x"] for q in pts], [q["y"] for q in pts]
                c = PALETTE[i % len(PALETTE)]
                if a.kind == "line":
                    ax.plot(xs, ys, "-o", ms=3, lw=1.6, color=c, label=_t(s.get("name")))
                else:
                    ax.scatter(xs, ys, s=16, color=c, alpha=0.85, label=_t(s.get("name")))
            if p.get("diagonal"):
                vals = [v for s in series for q in s.get("points") or [] for v in (q["x"], q["y"])]
                if vals:
                    ax.plot([min(vals), max(vals)], [min(vals), max(vals)], "--", color=COLOR["ink-4"], lw=1)
            ax.set_xlabel(_t(p.get("x", "")))
            ax.set_ylabel(_t(p.get("y", "")))
        elif a.kind == "bar":
            cats = p.get("categories") or []
            n = max(1, len(series))
            for i, s in enumerate(series):
                xs = [j + i / (n + 1) for j in range(len(cats))]
                ax.bar(xs, s.get("values") or [], width=1 / (n + 1), color=PALETTE[i % len(PALETTE)], label=_t(s.get("name")))
            _set_xticks(fig, ax, [j + (n - 1) / (2 * (n + 1)) for j in range(len(cats))], cats)
        elif a.kind == "parallel":
            dims, rows = p.get("dimensions") or [], p.get("rows") or []
            if not dims or not rows:
                return None
            cols = list(zip(*rows))
            norm = [[(v - min(c)) / ((max(c) - min(c)) or 1) for v in c] for c in cols]
            for r in zip(*norm):
                ax.plot(range(len(dims)), r, color=PALETTE[0], alpha=0.25, lw=1)
            _set_xticks(fig, ax, list(range(len(dims))), dims)
        else:
            return None
        if len(series) > 1:
            # 틀이 없으면 격자선이 범례 글자 위로 지나간다
            ax.legend(fontsize=8, frameon=True, framealpha=0.92, facecolor=COLOR["print-bg"], edgecolor=COLOR["line"], labelcolor=COLOR["ink-2"])
        ax.grid(color=COLOR["line"], alpha=0.7)
        ax.set_axisbelow(True)  # 격자가 막대 위로 지나가지 않게
        ax.tick_params(colors=COLOR["ink-3"])
        ax.xaxis.label.set_color(COLOR["ink-3"])
        ax.yaxis.label.set_color(COLOR["ink-3"])
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        for side in ("left", "bottom"):
            ax.spines[side].set_color(COLOR["line"])
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        return buf.getvalue()
    finally:
        plt.close(fig)


def image_bytes(project_id: int, a: models.Artifact) -> bytes | None:
    rel = (a.payload or {}).get("path")
    if not rel:
        return None
    f = (settings.files_dir / f"project_{project_id}" / rel).resolve()
    root = (settings.files_dir / f"project_{project_id}").resolve()
    return f.read_bytes() if root in f.parents and f.is_file() else None


def picture(project_id: int, a: models.Artifact) -> bytes | None:
    """결과의 그림. 결과 하나를 못 그린다고 문서 전체가 나가지 못하면 안 되므로 실패는 그 그림만의 안내로 끝낸다."""
    try:
        return image_bytes(project_id, a) if a.kind == "image" else chart_png(a)
    except Exception:  # noqa: BLE001 — AI·솔버가 준 자료의 모양은 믿을 수 없다
        log.warning("보고서 그림을 만들지 못했습니다: artifact %s", a.id, exc_info=True)
        return None


def _no_picture(a: models.Artifact) -> str:
    return NO_IMAGE if a.kind == "image" else NO_CHART


def _ordered(arts: list[models.Artifact]) -> list[models.Artifact]:
    return sorted(arts, key=lambda a: (KIND_ORDER.get(a.kind, 9), a.created_at))


# ---------- 작업 흐름 단계 번호 (웹 화면 orderNodes와 같은 규칙) ----------
def order_nodes(nodes: list[dict], edges: list[dict]) -> list[dict]:
    """실행 순서(엣지 기준 위상 정렬, 같은 순위는 원래 순서). 이 순서가 '1단계, 2단계' 번호다."""
    indeg = {n["id"]: 0 for n in nodes}
    for e in edges:
        indeg[e["target"]] = indeg.get(e["target"], 0) + 1
    left, out = list(nodes), []
    while left:
        i = next((k for k, n in enumerate(left) if indeg.get(n["id"], 0) == 0), 0)
        n = left.pop(i)
        out.append(n)
        for e in edges:
            if e["source"] == n["id"]:
                indeg[e["target"]] = indeg.get(e["target"], 1) - 1
    return out


def step_names(flow: models.Flow) -> dict[str, str]:
    """노드 id → '2단계 · DOE 60케이스'"""
    return {n["id"]: f"{i}단계 · {n.get('label', '')}" for i, n in enumerate(order_nodes(flow.nodes or [], flow.edges or []), 1)}


# ---------- AI 보고서 본문: markdown 사이에 ![제목](artifact:<id>) 줄로 결과를 넣는다 ----------
EMBED = re.compile(r"^\s*!\[(.*)\]\(artifact:(\d+)\)\s*$")  # 캡션에 [ ]가 들어가도 한 줄 전체로 본다


def embedded_ids(body: str) -> list[int]:
    return [p[2] for p in split_body(body) if p[0] == "art"]


def split_body(body: str) -> list[tuple[str, str, int | None]]:
    """본문을 ('md', 글, None) · ('art', 캡션, id) 조각으로 나눈다. 코드 블록 안의 삽입 줄은 예시이므로 글로 둔다."""
    parts: list[tuple[str, str, int | None]] = []
    buf: list[str] = []
    for line, state in fence_states(lines_of(body)):
        m = EMBED.match(line) if state == "text" else None
        if m:
            if buf:
                parts.append(("md", "\n".join(buf), None))
                buf = []
            parts.append(("art", m.group(1).strip(), int(m.group(2))))
        else:
            buf.append(line)
    if buf:
        parts.append(("md", "\n".join(buf), None))
    return parts


def _title_key(text: str) -> str:
    return " ".join(re.sub(r"[*_`]", "", text).split()).casefold()


HEADING = re.compile(r"^ {0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_OWN_CONTENTS = {"목차", "차례", "contents", "tableofcontents"}


def _drop_title(lines: list[str], title: str) -> list[str]:
    """첫 줄의 '# 제목'은 표지에 이미 있다. 제목과 달라도 뒤가 '##' 구역이면 문서 머리이므로 글로만 남기고 목차에 넣지 않는다."""
    i = next((k for k, line in enumerate(lines) if line.strip()), None)
    m = HEADING.match(lines[i]) if i is not None else None
    if not m or len(m[1]) != 1:
        return lines
    if _title_key(m[2]) == _title_key(title or ""):
        return lines[i + 1:]
    rest = [HEADING.match(line) for line, state in fence_states(lines[i + 1:]) if state == "text"]
    levels = {len(h[1]) for h in rest if h}
    return [*lines[:i], f"**{m[2]}**", *lines[i + 1:]] if 2 in levels and 1 not in levels else lines


def _drop_own_contents(lines: list[str]) -> list[str]:
    """옛 판·옛 Agent 본문이 직접 쓴 목차를 뺀다. 틀의 목차와 겹치고, 그 안의 링크는 어디로도 가지 않는다. 목록만 있는 구역일 때만 뺀다."""
    states = list(fence_states(lines))
    heads = [i for i, (line, state) in enumerate(states) if state == "text" and HEADING.match(line)]
    if not heads or _title_key(HEADING.match(lines[heads[0]])[2]).replace(" ", "") not in _OWN_CONTENTS:
        return lines
    start = heads[0]
    end = heads[1] if len(heads) > 1 else len(lines)
    section = [line for line in lines[start + 1:end] if line.strip()]
    if not section or any(states[k][1] != "text" for k in range(start, end)) or not all(report_md.is_list_line(line) for line in section):
        return lines
    return lines[:start] + lines[end:]


def clean_body(body: str, title: str) -> str:
    """틀이 이미 주는 것을 본문에서 뺀다(제목·직접 쓴 목차). AI 규칙은 '## 본문만'이지만 옛 판과 규칙을 어긴 글도 틀이 두 번 찍히면 안 된다."""
    lines = lines_of(body)
    kept = _drop_own_contents(_drop_title(lines, title))
    return body or "" if kept == lines else "\n".join(kept).lstrip("\n")


def _figure_label(kind: str, counters: dict[str, int]) -> str:
    key = "표" if kind == "table" else "그림"
    counters[key] = counters.get(key, 0) + 1
    return f"{key} {counters[key]}"


# ---------- 틀에 들어가는 발행 정보 ----------
@dataclass(frozen=True)
class Issue:
    doc_no: str
    edition: str
    issued: datetime
    author: str
    provenance: str
    colophon: str

    @property
    def issued_at(self) -> str:
        return f"{self.issued:%Y-%m-%d %H:%M}"


def make_issue(project: models.Project, rep: models.Report | None = None, edition: int = 1, requester: str = "") -> Issue:
    """표지·머리글·맺음말의 글. 사실만 적는다: AI가 썼고, 사람의 검토가 아직 필요하다."""
    who = f" (요청: {requester})" if requester else ""
    if rep:
        return Issue(f"DDS-P{project.id:04d}-R{rep.id:03d}", f"{edition}판", rep.created_at, f"AI{who}",
                     f"{ISSUER} 발행 · AI 작성, 담당 엔지니어 검토 필요",
                     f"이 문서는 {ISSUER}가 발행했습니다. AI가 작업 흐름과 결과를 읽고 썼으므로, 쓰기 전에 담당 엔지니어의 검토가 필요합니다.")
    return Issue(f"DDS-P{project.id:04d}-AUTO", "자동 모음", datetime.now(), f"자동 모음{who}",
                 f"{ISSUER} 발행 · 게시된 결과 자동 모음, 담당 엔지니어 검토 필요",
                 f"이 문서는 {ISSUER}가 게시된 결과를 자동으로 모은 것입니다. 해석이 들어 있지 않으며, 쓰기 전에 담당 엔지니어의 검토가 필요합니다.")


@cache
def _asset(name: str) -> bytes:
    return (ASSETS / name).read_bytes()


@cache
def _logo_svg() -> str:
    """HTML 안에 넣는 로고. HTML5 안의 svg에는 xmlns가 필요 없고, 파일에 외부 주소가 한 글자도 없게 뺀다."""
    svg = re.sub(r"<!--.*?-->", "", _asset("dds-logo.svg").decode("utf-8"), flags=re.S)
    svg = svg.replace(' xmlns="http://www.w3.org/2000/svg"', "").replace("<svg ", '<svg class="dds-logo" ', 1)
    return re.sub(r">\s+<", "><", svg).strip()


@cache
def _margin_logo_uri() -> str:
    svg = re.sub(r"<!--.*?-->", "", _asset("dds-logo.svg").decode("utf-8"), flags=re.S).replace("<svg ", '<svg width="11" height="11" ', 1)
    return "data:image/svg+xml;base64," + base64.b64encode(re.sub(r">\s+<", "><", svg).strip().encode()).decode()


def _favicon_uri() -> str:
    return "data:image/svg+xml;base64," + base64.b64encode(_asset("favicon.svg")).decode()


# ---------- 내용을 조각으로: HTML과 Word가 같은 조각·같은 목차를 쓴다 ----------
@dataclass
class Block:
    kind: str  # md 글 · figure 번호 붙은 그림·표 · section 결과 모음의 구역 제목 · flow 작업 흐름 · result 모음의 결과 · note 안내
    parsed: report_md.Parsed | None = None
    art: models.Artifact | None = None
    text: str = ""
    anchor: str = ""
    flow: models.Flow | None = None


def _cut(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


@dataclass
class Composed:
    title: str
    project: str
    summary: str  # 표지의 한 줄(결과 모음은 프로젝트 설명)
    blocks: list[Block]
    toc: list[TocItem]


COVER_PROJECT, COVER_SUMMARY = 120, 240


def compose(project: models.Project, arts: list[models.Artifact], flow: models.Flow | None, rep: models.Report | None = None) -> Composed:
    toc = report_md.Toc()
    blocks: list[Block] = []
    if rep:
        by_id = {a.id: a for a in arts}
        counters: dict[str, int] = {}
        for kind, text, aid in split_body(clean_body(rep.body, rep.title)):
            if kind == "md":
                blocks.append(Block("md", parsed=report_md.parse(text, toc)))
            elif (a := by_id.get(aid)) is not None:  # 지워진 결과는 파일에서 빠지고 번호도 세지 않는다
                blocks.append(Block("figure", art=a, text=f"{_figure_label(a.kind, counters)}. {text or a.title}"))
    else:
        if flow:
            title = f"작업 흐름 — {flow.title}"
            blocks += [Block("section", text=title, anchor=toc.add(2, title)), Block("flow", flow=flow)]
        for a in _ordered(arts):
            blocks += [Block("section", text=a.title, anchor=toc.add(2, a.title)), Block("result", art=a)]
        if not blocks:
            blocks.append(Block("note", text="게시된 결과가 아직 없습니다."))
        if len(desc := safe_text(project.description or "")) > COVER_SUMMARY:
            blocks.insert(0, Block("note", text=desc))  # 표지에는 앞부분만 두고 전문은 본문 첫머리에 둔다
    # 결과 모음은 프로젝트 이름이 곧 제목이므로 표지에 두 번 쓰지 않는다. 표지 한 쪽에 발행 정보까지 들어가게 긴 글은 줄인다
    return Composed(safe_text(rep.title if rep else project.name), _cut(safe_text(project.name), COVER_PROJECT) if rep else "",
                    "" if rep else _cut(safe_text(project.description or ""), COVER_SUMMARY), blocks, toc.items)


def _flow_lines(flow: models.Flow):
    for n in order_nodes(flow.nodes or [], flow.edges or []):
        status = n.get("status")
        yield status if status in STATUS else "", n.get("label", ""), STATUS.get(status, ""), n.get("detail") or ""


# ---------- HTML ----------
_CSS = """*{box-sizing:border-box}html{-webkit-text-size-adjust:100%}
@media (prefers-reduced-motion:no-preference){html{scroll-behavior:smooth}}
body{margin:0;background:var(--panel);color:var(--ink);font:14px/1.75 "Pretendard Variable",Pretendard,"Malgun Gothic","Apple SD Gothic Neo","Noto Sans KR",sans-serif;word-break:keep-all;overflow-wrap:anywhere}
.mark{display:inline-flex;align-items:center;gap:6px;font-weight:500;color:var(--ink);white-space:nowrap}
.mark b{font-weight:800;color:var(--logo-text)}.dds-logo{display:block;flex:none;width:1.5em;height:1.5em}
.doc-head{position:sticky;top:0;z-index:2;display:flex;align-items:center;justify-content:space-between;gap:16px;padding:8px 24px;background:var(--panel);border-bottom:1px solid var(--line);font-size:13px}
.doc-no{font-family:"Cascadia Mono",Consolas,"D2Coding",ui-monospace,monospace;font-size:12px;color:var(--ink-3);white-space:nowrap}
.sheet{max-width:1064px;margin:0 auto;padding:0 24px}
.cover{padding:48px 0 32px;border-bottom:1px solid var(--ink)}.cover .mark{font-size:20px;gap:8px}
h1{margin:32px 0 8px;font-size:40px;line-height:1.2;font-weight:650;letter-spacing:-.03em;text-wrap:balance}
.project{margin:0;font-size:16px;color:var(--ink-2)}.summary{margin:4px 0 0;color:var(--ink-3)}
.cover.long h1{font-size:28px}.cover.xlong h1{font-size:22px}.print-split{display:none}.print-split>p{margin:0}
.issue{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px 24px;margin:32px 0 0}.issue dt{font-size:12px;color:var(--ink-3)}.issue dd{margin:0;font-size:13px;font-weight:500}
.provenance{margin:24px 0 0;padding-left:10px;border-left:2px solid var(--brand);font-size:13px;color:var(--ink-2)}
.layout{display:grid;grid-template-columns:240px minmax(0,1fr);gap:48px;margin-top:32px}.layout.solo{display:block;max-width:776px}
.toc{position:sticky;top:56px;align-self:start;max-height:calc(100vh - 72px);overflow:auto;font-size:13px}
.toc-title{margin:0 0 8px;font-size:12px;font-weight:600;letter-spacing:.06em;color:var(--ink-3)}
.toc ol{list-style:none;margin:0;padding:0}.toc a{display:block;padding:4px 8px 4px 10px;border-left:2px solid var(--line-2);color:var(--ink-2);line-height:1.4;text-decoration:none}
.toc .l3 a{padding-left:22px;font-size:12px;color:var(--ink-3)}.toc a:hover{color:var(--ink);border-left-color:var(--ink-4)}
.toc a[aria-current]{color:var(--brand);border-left-color:var(--brand);font-weight:600}a:focus-visible{outline:2px solid var(--brand);outline-offset:2px}
main{min-width:0}main a{color:var(--brand)}
h2,h3{scroll-margin-top:56px}h2{margin:32px 0 8px;padding-top:12px;border-top:1px solid var(--line);font-size:20px;line-height:1.35;font-weight:650;letter-spacing:-.025em}main>h2:first-child{margin-top:0;padding-top:0;border-top:none}
h3{margin:24px 0 6px;font-size:16px;line-height:1.4;font-weight:650}h4,h5,h6{margin:16px 0 4px;font-size:14px;font-weight:650}
p,ul,ol{margin:0 0 10px}ul,ol{padding-left:22px}li{margin:2px 0}blockquote{margin:12px 0;padding-left:12px;border-left:2px solid var(--line);color:var(--ink-2)}hr{border:none;border-top:1px solid var(--line);margin:24px 0}
code{font-family:"Cascadia Mono",Consolas,"D2Coding",ui-monospace,monospace;font-size:.9em;background:var(--panel-2);border-radius:2px;padding:0 4px}
pre{margin:12px 0;padding:10px 12px;background:var(--panel-2);border-radius:6px;overflow-x:auto;font-size:12px;line-height:1.6;word-break:normal;overflow-wrap:normal}pre code{background:none;padding:0;font-size:inherit}
.tw{overflow-x:auto;margin:8px 0}table{border-collapse:collapse;width:100%;font-size:13px}th,td{border:1px solid var(--line);padding:6px 10px;word-break:normal;overflow-wrap:normal}
th{background:var(--panel-2);color:var(--ink-2);text-align:left;font-weight:600}td.r{text-align:right;font-variant-numeric:tabular-nums}
img{display:block;max-width:100%;height:auto;background:var(--print-bg);border:1px solid var(--line-2);border-radius:6px}
figure{margin:24px 0}figcaption{margin-top:8px;padding-top:8px;border-top:1px solid var(--line);font-size:12px;color:var(--ink-3)}
.meta{color:var(--ink-3);font-size:13px}.flow .st{margin-left:6px;font-size:12px;color:var(--ink-3)}.flow li.complete .st{color:var(--ok-ink)}.flow li.failed .st{color:var(--err-ink)}.flow li.in_progress .st{color:var(--warn-ink)}
.colophon{display:flex;gap:12px;margin-top:48px;padding:24px 0;border-top:1px solid var(--ink);font-size:13px;color:var(--ink-2)}.colophon .dds-logo{width:32px;height:32px}.colophon p{margin:0}.colophon .small{margin-top:4px;font-size:12px;color:var(--ink-3)}
.doc-foot{display:flex;justify-content:space-between;gap:16px;flex-wrap:wrap;padding:12px 24px;border-top:1px solid var(--line);font-size:12px;color:var(--ink-3)}
@media (max-width:960px){.layout{display:block}.toc{position:static;max-height:none;margin-bottom:24px;padding:12px;border:1px solid var(--line);border-radius:6px}.issue{grid-template-columns:repeat(2,minmax(0,1fr))}h1{font-size:28px}.cover{padding-top:32px}}
@media (max-width:480px){.doc-head,.doc-foot{padding-inline:16px}.sheet{padding:0 16px}}
@media print{@page wide{size:A4 landscape}@page{size:A4;margin:22mm 18mm 20mm;@top-left{content:url("@@LOGO_URI@@") " DDS Conversa 발행";font-size:8.5pt;color:var(--ink-3)}@top-right{content:"@@DOC_NO@@";font-size:8.5pt;color:var(--ink-3)}@bottom-left{content:"@@DOC_NO@@ · DDS Conversa";font-size:8.5pt;color:var(--ink-3)}@bottom-right{content:counter(page) " / " counter(pages);font-size:8.5pt;color:var(--ink-3)}}
@page:first{@top-left{content:none}@top-right{content:none}}
html{scroll-behavior:auto}body{background:var(--print-bg);font-size:10.5pt}.doc-head,.doc-foot{display:none}.sheet{max-width:none;padding:0}
.cover{padding:60mm 0 0;border-bottom:none;break-after:page}h1{font-size:28pt}.layout,.layout.solo{display:block;max-width:none;margin:0}
.toc{position:static;max-height:none;overflow:visible;margin:0;padding:0;border:none;break-after:page;font-size:10.5pt}.toc a,.toc a[aria-current]{padding:3px 0;border:none;color:var(--ink);font-weight:400}.toc .l3 a{padding-left:14px;color:var(--ink-2);font-size:10pt}.toc-title{font-size:14pt;color:var(--ink);letter-spacing:0}
main h2:not(:first-child){margin-top:20pt}h2,h3,h4{break-after:avoid}figure,pre,tr,img{break-inside:avoid}main a{color:inherit}.colophon{break-inside:avoid}.cover.long h1{font-size:20pt}.cover.xlong h1{font-size:15pt}
img{width:auto;max-height:200mm}pre{white-space:pre-wrap;overflow:visible;overflow-wrap:anywhere}
.tw{overflow:visible}table{font-size:9pt}th,td{padding:3pt 5pt;overflow-wrap:break-word}td:not(.r,.n){overflow-wrap:anywhere}td.r,td.n{white-space:nowrap}
.screen-part{display:none}.print-split{display:block}.print-split>p{margin:8pt 0 2pt}
.wide table{font-size:8pt}.wide.f7 table{font-size:7pt}.wide th,.wide td{padding:1.5pt 3pt}
.wide td.r,.wide td.n{white-space:normal;overflow-wrap:anywhere}.wide th{overflow-wrap:anywhere}
@supports (page:wide){main>figure.wide,main>.tw.wide,main>.print-split,main>h2:has(+figure.wide),main>h3:has(+figure.wide),main>h2:has(+.tw.wide),main>h3:has(+.tw.wide){page:wide}
.wide:not(.xwide) td.r,.wide:not(.xwide) td.n{white-space:nowrap;overflow-wrap:normal}.wide:not(.xwide) th{overflow-wrap:break-word}}}"""

_SCRIPT = """(function(){var links=[].slice.call(document.querySelectorAll('.toc a'));if(!links.length||!window.requestAnimationFrame)return;
var box=document.querySelector('.toc'),heads=links.map(function(a){return document.getElementById(a.getAttribute('href').slice(1));}),pin=null,busy=false,cur=-1,calm=false,timer=0;
function wait(){clearTimeout(timer);timer=setTimeout(function(){calm=true;},200);}
function show(i){if(i===cur)return;if(links[cur])links[cur].removeAttribute('aria-current');cur=i;var a=links[i];a.setAttribute('aria-current','true');
if(box.scrollHeight>box.clientHeight&&getComputedStyle(box).position==='sticky'){var t=a.offsetTop;if(t<box.scrollTop||t+a.offsetHeight>box.scrollTop+box.clientHeight)box.scrollTop=t-box.clientHeight/2;}}
function update(){busy=false;var i=0;if(pin!==null){show(pin);return;}for(var k=0;k<heads.length;k++){if(heads[k]&&heads[k].getBoundingClientRect().top<=72)i=k;}
if(window.innerHeight+window.scrollY>=document.documentElement.scrollHeight-2)i=heads.length-1;show(i);}
function ask(){if(!busy){busy=true;requestAnimationFrame(update);}}
links.forEach(function(a,i){a.addEventListener('click',function(){pin=i;calm=false;wait();ask();});});
['wheel','touchstart','keydown','pointerdown'].forEach(function(e){window.addEventListener(e,function(){pin=null;},{passive:true});});
window.addEventListener('scroll',function(){if(pin!==null){if(calm)pin=null;else wait();}ask();},{passive:true});window.addEventListener('resize',ask);update();})();"""


def _e(text) -> str:
    return html.escape(safe_text(str(text if text is not None else "")))


def _root_vars() -> str:
    return ":root{color-scheme:light;" + "".join(f"--{k}:{v};" for k, v in COLOR.items()) + "}"


def _cell(row: list, c: int) -> str:
    v = row[c] if c < len(row) else None
    text = _fmt(v)
    cls = "r" if isinstance(v, (int, float)) else "n" if report_md.is_number(text) else ""
    return f"<td class='{cls}'>{_e(text)}</td>" if cls else f"<td>{_e(text)}</td>"


def _one_table(cols: list, rows: list[list], picks, cls: str = "") -> str:
    # thead가 있어야 인쇄할 때 쪽마다 열 이름이 다시 나온다
    head = "".join(f"<th>{_e(_fmt(cols[c] if c < len(cols) else None))}</th>" for c in picks)
    body = "".join("<tr>" + "".join(_cell(r, c) for c in picks) + "</tr>" for r in rows)
    attr = f' class="{cls}"' if cls else ""
    return f"<table{attr}><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _table_html(p: dict) -> str:
    cols, rows = p.get("columns", []), p.get("rows", [])[:200]
    plan = plan_table([[_fmt(c) for c in cols], *[[_fmt(v) for v in r] for r in rows]])
    n = max((len(cols), *(len(r) for r in rows)), default=0)
    if len(plan.groups) == 1:
        return _one_table(cols, rows, range(n), plan.css)
    # 화면은 가로 스크롤 표 하나, 인쇄는 값이 끊기지 않게 첫 열을 반복해 나눈 표들
    css = plan.css.replace(" xwide", "")
    parts = "".join(f"<p class='meta'>표가 넓어 나눠 싣습니다({k}/{len(plan.groups)}). 첫 열은 반복입니다.</p>" + _one_table(cols, rows, g, css)
                    for k, g in enumerate(plan.groups, 1))
    return _one_table(cols, rows, range(n), f"{css} screen-part") + f'<div class="print-split">{parts}</div>'


def _art_html(project_id: int, a: models.Artifact) -> str:
    p = a.payload or {}
    if a.kind == "markdown":
        return report_md.parse(p.get("text", ""), shift=1).html  # 결과 안의 제목은 구역 아래 단계이고 목차에 넣지 않는다
    if a.kind == "table":
        return _table_html(p)
    data = picture(project_id, a)
    if not data:
        return f"<p class='meta'>{_no_picture(a)}</p>"
    # 올릴 수 있는 형식 중 svg만 브라우저가 내용으로 알아보지 못하므로 종류를 밝힌다(img 안의 svg는 스크립트·외부 요청이 돌지 않는다)
    mime = "svg+xml" if b"<svg" in data[:1024] else "png"
    return f"<img src='data:image/{mime};base64,{base64.b64encode(data).decode()}' alt='{_e(a.title)}'>"


def _block_html(project_id: int, b: Block) -> str:
    if b.kind == "md":
        return b.parsed.html
    if b.kind == "figure":
        inner = _art_html(project_id, b.art)
        attr = ' class="wide"' if '<table class="wide' in inner else ""  # 설명도 표와 같은 가로 쪽에 있어야 한다
        return f"<figure{attr}>{inner}<figcaption>{_e(b.text)}</figcaption></figure>"
    if b.kind == "section":
        return f'<h2 id="{b.anchor}">{_e(b.text)}</h2>'
    if b.kind == "flow":
        items = "".join(f"<li class='{cls}'>{_e(label)}{' — ' + _e(detail) if detail else ''}{f'<span class=st>{word}</span>' if word else ''}</li>"
                        for cls, label, word, detail in _flow_lines(b.flow))
        goal = f"<p class='meta'>{_e(b.flow.goal)}</p>" if b.flow.goal else ""
        return f"{goal}<ol class='flow'>{items}</ol>"
    if b.kind == "result":
        caption = (b.art.payload or {}).get("caption") if b.art.kind == "image" else None
        return _art_html(project_id, b.art) + (f"<p class='meta'>{_e(caption)}</p>" if caption else "")
    return f"<p class='meta'>{_e(b.text)}</p>"


def _toc_html(items: list[TocItem]) -> str:
    rows = "".join(f'<li class="l{t.level}"><a href="#{t.anchor}">{_e(t.text)}</a></li>' for t in items)
    return f'<nav class="toc" aria-label="목차"><p class="toc-title">목차</p><ol>{rows}</ol></nav>'


def _cover_class(title: str) -> str:
    return "cover xlong" if len(title) > 120 else "cover long" if len(title) > 60 else "cover"


def build_html(project: models.Project, arts: list[models.Artifact], flow: models.Flow | None,
               rep: models.Report | None = None, issue: Issue | None = None) -> str:
    """파일 하나로 끝나는 문서: 파비콘·로고·그림을 안에 넣어 외부 요청이 없다. 목차는 스크립트 없이도 동작하고, 스크립트는 현재 위치 표시만 한다."""
    d, i = compose(project, arts, flow, rep), issue or make_issue(project, rep)
    logo = _logo_svg()
    wordmark = f'<span class="mark">{logo}<b>DDS</b><span>Conversa</span></span>'
    body = "".join(_block_html(project.id, b) for b in d.blocks)
    # 넓은 표가 화면 폭을 밀어내지 않게 표마다 가로 스크롤 상자를 씌운다(코드 안의 <table>은 이미 글자로 바뀌어 있다)
    body = re.sub(r'<table(?: class="([^"]*)")?>', lambda m: f'<div class="tw{" " + m[1] if m[1] else ""}"><table>', body).replace("</table>", "</table></div>")
    facts = (("문서 번호", i.doc_no), ("판", i.edition), ("발행 일시", i.issued_at), ("작성자", i.author))
    dl = "".join(f"<div><dt>{k}</dt><dd>{_e(v)}</dd></div>" for k, v in facts)
    # 문서 번호는 숫자로 만든 값이라 CSS에 넣어도 안전하다. 머리글 로고는 크기가 있어야 쪽 여백에 맞는다
    css = _root_vars() + _CSS.replace("@@DOC_NO@@", i.doc_no).replace("@@LOGO_URI@@", _margin_logo_uri())
    sub = (f'<p class="project">{_e(d.project)}</p>' if d.project else "") + (f'<p class="summary">{_e(d.summary)}</p>' if d.summary else "")
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="color-scheme" content="light"><meta name="generator" content="{ISSUER}"><meta name="author" content="{ISSUER}">
<title>{_e(d.title)} · {ISSUER}</title><link rel="icon" type="image/svg+xml" href="{_favicon_uri()}"><style>{css}</style></head><body>
<header class="doc-head">{wordmark}<span class="doc-no">{i.doc_no}</span></header>
<div class="sheet"><section class="{_cover_class(d.title)}" aria-label="표지">{wordmark}<h1>{_e(d.title)}</h1>{sub}<dl class="issue">{dl}</dl><p class="provenance">{_e(i.provenance)}</p></section>
<div class="layout{'' if d.toc else ' solo'}">{_toc_html(d.toc) if d.toc else ''}<main>{body}</main></div>
<section class="colophon" aria-label="맺음말">{logo}<div><p>{_e(i.colophon)}</p><p class="small">{i.doc_no} · {_e(i.edition)} · {i.issued_at}</p></div></section></div>
<footer class="doc-foot"><span>{ISSUER} 발행</span><span class="doc-no">{i.doc_no}</span></footer>{f'<script>{_SCRIPT}</script>' if d.toc else ''}</body></html>"""


# ---------- Word ----------
FONT, MONO = "Malgun Gothic", "Consolas"
PORTRAIT_MM, LANDSCAPE_MM = 160.0, 247.0  # 글자 폭: A4 - 좌우 여백 25씩
IMAGE_MAX_MM = 200.0  # 글자 높이(250) 안에 제목·설명 자리를 남긴다. 넘치면 그림이 꼬리말을 덮는다
# 문단 속성의 순서가 어긋나면 Word가 파일을 고치겠다고 묻는다. python-docx가 모르는 속성은 이 순서로 끼워 넣는다
_PPR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
              "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
              "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
              "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
              "w:rPr", "w:sectPr", "w:pPrChange")
_BLOCK_TAGS = {"p", "ul", "ol", "pre", "blockquote", "table", "hr", "h1", "h2", "h3", "h4", "h5", "h6", "div"}


def _hex(name: str) -> str:
    return COLOR[name].lstrip("#").upper()


def _rgb(name: str) -> RGBColor:
    return RGBColor.from_string(_hex(name))


def _ppr_add(par, el) -> None:
    tag = el.tag.split("}")[1]
    after = _PPR_ORDER[_PPR_ORDER.index(f"w:{tag}") + 1:]
    par._p.get_or_add_pPr().insert_element_before(el, *after)


def _border(par, side: str, color: str, size: int = 6, space: int = 6) -> None:
    box = par._p.get_or_add_pPr().find(qn("w:pBdr"))
    if box is None:
        box = OxmlElement("w:pBdr")
        _ppr_add(par, box)
    el = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", _hex(color))):
        el.set(qn(f"w:{k}"), v)
    box.append(el)


def _style_font(style, size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    style.font.name = FONT
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    for attr in list(fonts.attrib):  # 기본 틀의 테마 글꼴이 남으면 한글이 다른 글꼴로 나온다
        del fonts.attrib[attr]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT)
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = _rgb(color)
    if bold is not None:
        style.font.bold = bold


def _setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    _style_font(normal, 10, "ink")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15  # 맑은 고딕은 글꼴 자체의 줄 높이가 커서 이 값이 HTML의 줄 간격과 비슷하다
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")
    normal.element.get_or_add_rPr().append(lang)
    # 한글과 영문·숫자 사이를 Word가 저절로 띄우면 '2 판'·'AI 가'처럼 보인다. 쓴 그대로 둔다
    ppr = normal.element.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        off = OxmlElement(tag)
        off.set(qn("w:val"), "0")
        ppr.insert_element_before(off, *_PPR_ORDER[_PPR_ORDER.index(tag) + 1:])
    # 기본 틀은 14라서 Word가 '호환 모드'로 열고 표 왼쪽 선이 여백 밖으로 나간다
    for s in doc.settings.element.iter(qn("w:compatSetting")):
        if s.get(qn("w:name")) == "compatibilityMode":
            s.set(qn("w:val"), "15")
    for name, size, before, after in (("Heading 1", 15, 18, 6), ("Heading 2", 12, 14, 4), ("Heading 3", 10.5, 10, 3), ("Heading 4", 10, 8, 2)):
        st = doc.styles[name]
        _style_font(st, size, "ink", True)
        st.font.italic = False
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.0
    title = doc.styles["Title"]
    _style_font(title, 26, "ink", True)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.line_spacing = 1.0
    box = title.element.get_or_add_pPr().find(qn("w:pBdr"))
    if box is not None:  # 기본 틀 제목의 파란 밑줄
        title.element.pPr.remove(box)
    _style_font(doc.styles["Caption"], 9, "ink-3", False)
    for name in ("Header", "Footer"):
        _style_font(doc.styles[name], 8.5, "ink-3")
        doc.styles[name].paragraph_format.tab_stops.clear_all()  # 편지지 폭 기준의 기본 탭을 지운다
        doc.styles[name].paragraph_format.space_after = Pt(0)


def _field(par, code: str) -> None:
    """쪽 번호 필드. 머리글·꼬리말의 쪽 번호는 Word가 스스로 계산하므로 '필드 업데이트' 질문이 뜨지 않는다."""
    for kind, text in (("begin", None), (None, code), ("separate", None), ("text", "1"), ("end", None)):
        run = par.add_run()
        if kind in ("begin", "separate", "end"):
            ch = OxmlElement("w:fldChar")
            ch.set(qn("w:fldCharType"), kind)
            run._r.append(ch)
        elif kind == "text":
            run.text = text
        else:
            ins = OxmlElement("w:instrText")
            ins.set(qn("xml:space"), "preserve")
            ins.text = f" {text} "
            run._r.append(ins)


def _foot(par, i: Issue, width: float) -> None:
    par.paragraph_format.tab_stops.add_tab_stop(Mm(width), WD_TAB_ALIGNMENT.RIGHT)
    par.add_run(f"{i.doc_no} · {ISSUER} 발행\t")
    _field(par, "PAGE")
    par.add_run(" / ")
    _field(par, "NUMPAGES")
    par.add_run(" 쪽")


def _frame(sec, d: Composed, i: Issue, width: float) -> None:
    """모든 쪽의 머리글(로고·발행처·제목)과 꼬리말. 가로 쪽은 폭이 달라 탭 위치가 다르므로 절마다 새로 만든다."""
    sec.header.is_linked_to_previous = sec.footer.is_linked_to_previous = False
    head = sec.header.paragraphs[0]
    head.paragraph_format.tab_stops.add_tab_stop(Mm(width), WD_TAB_ALIGNMENT.RIGHT)
    _center_line(head)
    head.add_run().add_picture(io.BytesIO(_asset("dds-logo.png")), height=Mm(4.5))
    name = head.add_run(f" {ISSUER}")
    name.bold, name.font.color.rgb = True, _rgb("logo-text")
    short = d.title if len(d.title) <= 34 else d.title[:33] + "…"  # 긴 제목은 머리글이 두 줄이 되지 않게 줄인다
    head.add_run(f" 발행\t{short}")
    _border(head, "bottom", "line", 4, 4)
    _foot(sec.footer.paragraphs[0], i, width)


def _page_size(sec, landscape: bool) -> None:
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin, sec.bottom_margin = Mm(25), Mm(22)
    sec.header_distance, sec.footer_distance = Mm(11), Mm(10)


def _setup_page(doc, d: Composed, i: Issue) -> None:
    sec = doc.sections[0]
    _page_size(sec, False)
    sec.different_first_page_header_footer = True  # 표지는 머리글 없이 큰 로고로 보인다. 꼬리말은 표지에도 있다
    _frame(sec, d, i, PORTRAIT_MM)
    _foot(sec.first_page_footer.paragraphs[0], i, PORTRAIT_MM)


def _center_line(par) -> None:
    """로고와 글자를 줄 가운데에 맞춘다."""
    ta = OxmlElement("w:textAlignment")
    ta.set(qn("w:val"), "center")
    _ppr_add(par, ta)


def _wordmark(par, logo_mm: float, size: float) -> None:
    _center_line(par)
    par.add_run().add_picture(io.BytesIO(_asset("dds-logo.png")), height=Mm(logo_mm))
    a = par.add_run(" DDS")
    a.bold, a.font.size, a.font.color.rgb = True, Pt(size), _rgb("logo-text")
    b = par.add_run(" Conversa")
    b.font.size = Pt(size)


def _table_lines(tbl, *sides: str) -> None:
    box = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        on = side in sides
        for k, v in (("val", "single" if on else "nil"), ("sz", "4"), ("space", "0"), ("color", _hex("line"))):
            el.set(qn(f"w:{k}"), v)
        box.append(el)
    tbl._tbl.tblPr.insert_element_before(box, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def _cover(doc, d: Composed, i: Issue) -> None:
    _wordmark(doc.add_paragraph(), 16, 18)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(90 if len(d.title) <= 70 else 40)
    title = doc.add_paragraph(d.title, style="Title")
    # 제목이 길어도 발행 정보가 표지 한 쪽에 남아야 한다. 글자 수로 크기를 고른다
    title.runs[0].font.size = Pt(26 if len(d.title) <= 40 else 22 if len(d.title) <= 70 else 18 if len(d.title) <= 120 else 15)
    if d.project:
        sub = doc.add_paragraph(d.project)
        sub.runs[0].font.size, sub.runs[0].font.color.rgb = Pt(13), _rgb("ink-2")
    if d.summary:
        doc.add_paragraph(d.summary).runs[0].font.color.rgb = _rgb("ink-3")
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(36)
    facts = (("문서 번호", i.doc_no), ("판", i.edition), ("발행 일시", i.issued_at), ("작성자", i.author))
    tbl = doc.add_table(rows=len(facts), cols=2)
    tbl.alignment, tbl.autofit = WD_TABLE_ALIGNMENT.LEFT, False
    _table_lines(tbl, "top", "bottom", "insideH")
    for row, (k, v) in zip(tbl.rows, facts):
        row.cells[0].width, row.cells[1].width = Mm(32), Mm(128)
        for cell, text, color in ((row.cells[0], k, "ink-3"), (row.cells[1], v, "ink")):
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(3)
            par.add_run(safe_text(text)).font.color.rgb = _rgb(color)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)  # 왼쪽 선이 띄움까지 길어지지 않게 띄움은 빈 문단으로 둔다
    note = doc.add_paragraph(i.provenance)
    note.paragraph_format.left_indent = Mm(3)
    note.runs[0].font.color.rgb = _rgb("ink-2")
    _border(note, "left", "brand", 12, 6)
    doc.add_page_break()


def _contents(doc, items: list[TocItem]) -> None:
    """누르면 바로 가는 목차: 책갈피 + w:hyperlink. 필드(TOC)가 아니라서 여는 순간 '필드 업데이트' 질문 없이 그대로 동작한다.
    쪽 번호는 넣지 않는다: Word가 쪽을 나눠 봐야 알 수 있어 서버가 맞는 값을 쓸 수 없고, 틀린 번호보다 없는 편이 낫다(product-spec 3-5절)."""
    if not items:
        return
    head = doc.add_paragraph("목차")
    head.runs[0].bold, head.runs[0].font.size = True, Pt(15)
    hint = doc.add_paragraph("화면에서 볼 때는 Ctrl 키를 누른 채 제목을 누르면 그 위치로 갑니다.")
    hint.runs[0].font.size, hint.runs[0].font.color.rgb = Pt(9), _rgb("ink-3")
    hint.paragraph_format.space_after = Pt(10)
    for t in items:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4 if t.level == 2 else 2)
        if t.level == 3:
            par.paragraph_format.left_indent = Mm(6)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), t.anchor)
        link.set(qn("w:history"), "1")
        run = par.add_run(safe_text(t.text))
        run.font.color.rgb = _rgb("ink" if t.level == 2 else "ink-2")
        link.append(run._r)
        par._p.append(link)
    doc.add_page_break()


class _WordBody:
    """해석된 markdown 나무와 결과를 Word 문단·표로 옮긴다."""

    def __init__(self, doc, project_id: int, d: Composed, issue: Issue):
        self.doc, self.project_id, self.marks = doc, project_id, 0
        self.d, self.issue, self.wide = d, issue, False

    def orient(self, wide: bool) -> None:
        """넓은 표는 가로 쪽에 싣는다. 쪽 방향은 절(section) 단위이므로 절을 나눈다.
        제목은 방향을 바꾸지 않고 다음 내용을 따라간다: 끝에 남은 제목을 새 절로 옮겨, 앞 쪽에 제목만 남거나 넓은 표 사이에 빈 쪽이 생기지 않게 한다."""
        if wide == self.wide:
            return
        body = self.doc.element.body
        moving: list = []
        for el in reversed(list(body)[:-1]):
            style = el.find(f"{qn('w:pPr')}/{qn('w:pStyle')}") if el.tag == qn("w:p") else None
            if style is None or not style.get(qn("w:val")).startswith("Heading"):
                break
            moving.insert(0, el)
        sec = self.doc.add_section()
        sec.different_first_page_header_footer = False
        _page_size(sec, wide)
        _frame(sec, self.d, self.issue, LANDSCAPE_MM if wide else PORTRAIT_MM)
        mark = body.xpath("./w:p[w:pPr/w:sectPr]")[-1]
        Paragraph(mark, self.doc._body).paragraph_format.space_after = Pt(0)
        rpr = OxmlElement("w:rPr")  # 절을 끊는 빈 문단이 줄 하나를 차지하지 않게
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "2")
        rpr.append(size)
        mark.pPr.insert_element_before(rpr, "w:sectPr", "w:pPrChange")
        anchor = mark
        for el in moving:
            anchor.addnext(el)
            anchor = el
        self.wide = wide

    def blocks(self, root) -> None:
        for el in root:
            self.block(el)

    def block(self, el, indent: float = 0, quote: bool = False) -> None:
        is_heading = re.fullmatch(r"h[2-6]", el.tag)
        if el.tag != "table" and not is_heading:
            self.orient(False)
        if is_heading:
            self.heading(el)
        elif el.tag == "p":
            if el.text or len(el):
                par = self.doc.add_paragraph()
                if indent:
                    par.paragraph_format.left_indent = Mm(indent)
                if quote:
                    _border(par, "left", "line", 12, 6)
                self.inline(par, el)
        elif el.tag in ("ul", "ol"):
            self.list(el)
        elif el.tag == "pre":
            self.code(el)
        elif el.tag == "table":
            rows = [list(tr) for tr in el.iter("tr")]
            self.table(rows, header=bool(el.find("thead") is not None or (rows and rows[0] and rows[0][0].tag == "th")))
        elif el.tag == "blockquote":
            for ch in el:
                self.block(ch, indent + 4, True)
        elif el.tag == "hr":
            _border(self.doc.add_paragraph(), "bottom", "line", 4, 1)
        elif el.text or len(el):
            self.inline(self.doc.add_paragraph(), el)

    def heading(self, el) -> None:
        par = self.doc.add_heading("", level=min(4, int(el.tag[1]) - 1))
        self.inline(par, el)
        if el.get("id"):
            self.bookmark(par, el.get("id"))

    def section(self, text: str, anchor: str) -> None:
        par = self.doc.add_heading(safe_text(text), level=1)
        self.bookmark(par, anchor)

    def bookmark(self, par, name: str) -> None:
        self.marks += 1
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(self.marks))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(self.marks))
        pos = 1 if par._p.pPr is not None else 0
        par._p.insert(pos, start)
        par._p.append(end)

    def inline(self, par, el, bold=False, italic=False, code=False) -> None:
        self.text(par, el.text, bold, italic, code)
        for ch in el:
            if ch.tag in _BLOCK_TAGS:
                continue
            if ch.tag == "br":
                par.add_run().add_break()
            elif ch.tag == "a" and ch.get("href"):
                self.link(par, ch)
            else:
                self.inline(par, ch, bold or ch.tag in ("strong", "b"), italic or ch.tag in ("em", "i"), code or ch.tag == "code")
            self.text(par, ch.tail, bold, italic, code)

    def text(self, par, text: str | None, bold=False, italic=False, code=False) -> None:
        # 줄바꿈은 br로 이미 나뉘어 있다. 남은 \\n은 해석기가 보기 좋게 넣은 것이다
        text = report_md.unescape(text or "" if code else (text or "").replace("\n", ""))
        if not text:
            return
        run = par.add_run(safe_text(text))
        run.bold, run.italic = bold or None, italic or None
        if code:
            run.font.name, run.font.color.rgb = MONO, _rgb("ink-2")

    def link(self, par, a) -> None:
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), par.part.relate_to(report_md.unescape(a.get("href")), RT.HYPERLINK, is_external=True))
        run = par.add_run(safe_text(plain(a)))
        run.font.color.rgb, run.underline = _rgb("brand"), True
        link.append(run._r)
        par._p.append(link)

    def new_num(self, style: str, start: int = 1) -> int:
        """번호 목록마다 1부터 다시 시작한다. 같은 스타일을 그냥 쓰면 문서 전체가 한 목록으로 이어진다."""
        numbering = self.doc.part.numbering_part.element
        base = self.doc.styles[style].element.pPr.numPr.numId.val
        num = numbering.add_num(numbering.num_having_numId(base).abstractNumId.val)
        num.add_lvlOverride(0).add_startOverride(start)
        return num.numId

    def list_par(self, style: str, num: int | None):
        par = self.doc.add_paragraph(style=style)
        if num is not None:
            numpr = par._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = num
        return par

    def list(self, el, depth: int = 0) -> None:
        base = "List Number" if el.tag == "ol" else "List Bullet"
        style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
        num = self.new_num(style, int(el.get("start") or 1)) if el.tag == "ol" else None
        for li in el.findall("li"):
            par = self.list_par(style, num)
            self.inline(par, li)
            used = bool(par.text)
            for ch in li:
                if ch.tag in ("ul", "ol"):
                    self.list(ch, depth + 1)
                elif ch.tag == "p" and not used:
                    self.inline(par, ch)
                    used = True
                elif ch.tag in _BLOCK_TAGS:
                    self.block(ch, 6 * (depth + 1))

    def code(self, el) -> None:
        par = self.doc.add_paragraph()
        shd = OxmlElement("w:shd")
        for k, v in (("val", "clear"), ("color", "auto"), ("fill", _hex("panel-2"))):
            shd.set(qn(f"w:{k}"), v)
        _ppr_add(par, shd)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_after = Pt(8)
        run = par.add_run(safe_text(report_md.unescape("".join(el.itertext())).rstrip("\n")))
        run.font.name, run.font.size = MONO, Pt(8.5)

    def table(self, rows: list[list], header: bool = True) -> None:
        """표 한 벌의 모양. 칸은 해석된 요소이거나 글자(숫자는 오른쪽 맞춤)다.
        좁은 칸에서 숫자가 '11.' '111'처럼 갈라지면 잘못 읽으므로 끊기지 않을 폭을 먼저 잡는다: 세로 쪽 → 가로 쪽·작은 글자 → 첫 열 반복해 나누기."""
        n = max((len(r) for r in rows), default=0)
        if not n:
            return
        texts = [[plain(v) if isinstance(v, etree.Element) else _fmt(v) for v in r] for r in rows]
        plan = plan_table(texts, PORTRAIT_MM, LANDSCAPE_MM)
        self.orient(plan.landscape)
        for k, cols in enumerate(plan.groups, 1):
            if len(plan.groups) > 1:
                note = self.doc.add_paragraph(f"표가 넓어 나눠 싣습니다({k}/{len(plan.groups)}). 첫 열은 반복입니다.")
                note.runs[0].font.size, note.runs[0].font.color.rgb = Pt(9), _rgb("ink-3")
                note.paragraph_format.keep_with_next = True
            self.one_table(rows, cols, plan, header)

    def one_table(self, rows: list[list], cols: tuple[int, ...], plan: TablePlan, header: bool) -> None:
        tbl = self.doc.add_table(rows=len(rows), cols=len(cols), style="Table Grid")
        tbl.autofit = False
        _table_lines(tbl, "top", "left", "bottom", "right", "insideH", "insideV")
        if plan.tight:
            mar = OxmlElement("w:tblCellMar")
            for side in ("left", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), "45")
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tbl._tbl.tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")
        widths = [Mm(w) for w in plan.widths(cols, LANDSCAPE_MM if plan.landscape else PORTRAIT_MM)]
        for grid, w in zip(tbl._tbl.tblGrid.gridCol_lst, widths):
            grid.w = w
        for r, (tr, cells) in enumerate(zip(tbl._tbl.tr_lst, rows)):
            is_head = r == 0 and header
            if is_head:
                tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # 쪽이 넘어가도 머리 줄을 다시 보인다
            # row.cells는 부를 때마다 표 전체를 훑어 큰 표에서 오래 걸린다. 칸 요소를 바로 쓴다
            for tc, c, w in zip(tr.tc_lst, cols, widths):
                tc.width = w
                cell = _Cell(tc, tbl)
                value = cells[c] if c < len(cells) else ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.line_spacing = 1.0
                if isinstance(value, etree.Element):
                    self.inline(par, value)
                    # markdown 표의 '---:' 맞춤을 살린다
                    side = re.search(r"text-align:\s*(right|center)", value.get("style") or "")
                    if side:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if side[1] == "right" else WD_ALIGN_PARAGRAPH.CENTER
                else:
                    par.add_run(safe_text(_fmt(value)))
                    if isinstance(value, (int, float)) and not is_head:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in par.runs:
                    run.font.size = Pt(plan.font)
                    if is_head:
                        run.bold = True
                if is_head:
                    shd = OxmlElement("w:shd")
                    for k, v in (("val", "clear"), ("color", "auto"), ("fill", _hex("panel-2"))):
                        shd.set(qn(f"w:{k}"), v)
                    tc.get_or_add_tcPr().append(shd)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def picture(self, a: models.Artifact) -> None:
        """그림은 원래 크기로, 쪽을 넘지 않게 줄여서만 넣는다. 작은 그림을 늘리면 흐려지고 긴 그림은 꼬리말을 덮는다."""
        self.orient(False)
        data = picture(self.project_id, a)
        info = None
        if data:
            try:
                info = DocxImage.from_blob(data)
            except Exception:  # noqa: BLE001 — Word가 못 읽는 형식이면 그 그림만 안내로 바꾼다
                log.warning("Word에 넣을 수 없는 그림입니다: artifact %s", a.id, exc_info=True)
        if info is None:
            par = self.doc.add_paragraph(NO_WORD_IMAGE if data else _no_picture(a))
            par.runs[0].font.color.rgb = _rgb("ink-3")
        else:
            w, h = info.px_width / 96, info.px_height / 96  # 인치. HTML 화면의 원래 크기와 같은 기준
            rate = min(1.0, PORTRAIT_MM / 25.4 / w, IMAGE_MAX_MM / 25.4 / h)
            self.doc.add_picture(io.BytesIO(data), width=Inches(w * rate), height=Inches(h * rate))
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True  # 그림과 설명이 쪽으로 갈라지지 않게

    def art(self, a: models.Artifact) -> None:
        p = a.payload or {}
        if a.kind == "markdown":
            self.blocks(report_md.parse(p.get("text") or "", shift=1).root)
        elif a.kind == "table":
            cols, rows = p.get("columns", []), p.get("rows", [])[:200]
            if cols:
                self.table([list(cols), *[list(r[: len(cols)]) for r in rows]])
        else:
            self.picture(a)

    def flow(self, flow: models.Flow) -> None:
        if flow.goal:
            self.doc.add_paragraph(safe_text(flow.goal))
        num = self.new_num("List Number")
        for _cls, label, word, detail in _flow_lines(flow):
            tail = f" ({word})" if word else ""
            self.list_par("List Number", num).add_run(safe_text(f"{label}{tail}{' — ' + detail if detail else ''}"))

    def put(self, b: Block) -> None:
        if b.kind == "md":
            self.blocks(b.parsed.root)
        elif b.kind == "figure":
            self.art(b.art)
            self.doc.add_paragraph(safe_text(b.text), style="Caption")
        elif b.kind == "section":
            self.section(b.text, b.anchor)
        elif b.kind == "flow":
            self.orient(False)
            self.flow(b.flow)
        elif b.kind == "result":
            self.art(b.art)
            caption = (b.art.payload or {}).get("caption") if b.art.kind == "image" else None
            if caption:
                self.doc.add_paragraph(safe_text(caption), style="Caption")
        else:
            self.orient(False)
            self.doc.add_paragraph(safe_text(b.text))


def _colophon(doc, i: Issue) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(30)
    par.paragraph_format.keep_with_next = True
    _border(par, "top", "ink", 6, 10)
    _wordmark(par, 6, 10)
    text = doc.add_paragraph(i.colophon)
    text.paragraph_format.keep_with_next = True
    text.runs[0].font.color.rgb = _rgb("ink-2")
    small = doc.add_paragraph(safe_text(f"{i.doc_no} · {i.edition} · {i.issued_at}"))
    small.runs[0].font.size, small.runs[0].font.color.rgb = Pt(9), _rgb("ink-3")


def _properties(doc, d: Composed, i: Issue) -> None:
    cp = doc.core_properties  # 기본 틀의 'python-docx' 표시를 지우고 발행처를 남긴다
    cp.author = cp.last_modified_by = ISSUER
    cp.title = d.title[:255]
    cp.subject = safe_text(f"{i.doc_no} · {d.project or d.title}")[:255]
    cp.comments = i.provenance
    cp.category = f"{ISSUER} 보고서"
    cp.keywords = ""
    cp.language = "ko-KR"
    cp.created = cp.modified = i.issued
    cp.revision = 1


def build_docx(project: models.Project, arts: list[models.Artifact], flow: models.Flow | None,
               rep: models.Report | None = None, issue: Issue | None = None) -> bytes:
    d, i = compose(project, arts, flow, rep), issue or make_issue(project, rep)
    doc = Document()
    _setup_styles(doc)
    _setup_page(doc, d, i)
    _cover(doc, d, i)
    _contents(doc, d.toc)
    body = _WordBody(doc, project.id, d, i)
    for b in d.blocks:
        body.put(b)
    body.orient(False)
    _colophon(doc, i)
    _properties(doc, d, i)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- AI가 보고서를 쓸 때 읽는 자료 ----------
def _num(v):
    return float(f"{v:.5g}") if isinstance(v, float) else v


def summarize(a: models.Artifact) -> dict:
    """결과 하나를 AI가 읽기 좋은 크기로 줄인다(표는 앞 12행, 그래프는 축·범위)."""
    p = a.payload or {}
    if a.kind == "table":
        rows = p.get("rows") or []
        return {"columns": p.get("columns", []), "rows": [[_num(v) for v in r] for r in rows[:12]], "n_rows": len(rows)}
    if a.kind in ("scatter", "line"):
        out = []
        for s in p.get("series") or []:
            pts = s.get("points") or []
            xs, ys = [q["x"] for q in pts], [q["y"] for q in pts]
            out.append({"name": s.get("name"), "n": len(pts),
                        "x_range": [_num(min(xs)), _num(max(xs))] if xs else None,
                        "y_range": [_num(min(ys)), _num(max(ys))] if ys else None, "y_last": _num(ys[-1]) if ys else None})
        return {"x": p.get("x"), "y": p.get("y"), "series": out}
    if a.kind == "bar":
        return {"categories": (p.get("categories") or [])[:30],
                "series": [{"name": s.get("name"), "values": [_num(v) for v in (s.get("values") or [])[:30]]} for s in p.get("series") or []]}
    if a.kind == "parallel":
        return {"dimensions": p.get("dimensions"), "n_rows": len(p.get("rows") or [])}
    if a.kind == "markdown":
        return {"text": (p.get("text") or "")[:4000]}
    return {"caption": p.get("caption", "")}


def _fmt(v) -> str:
    if isinstance(v, float):
        return f"{v:.6g}"
    return "" if v is None else str(v)
