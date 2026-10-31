"""보고서 틀(표지·목차·본문·맺음말): AI 보고서든 결과 모음이든, 내용이 무엇이든 틀은 그대로이고 DDS Conversa 발행임이 보여야 한다."""
import base64
import io
import re
import shutil
import subprocess
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from server.tests.test_api import _auth, client  # noqa: F401 — 서버 코드보다 먼저 불러야 임시 데이터 폴더를 쓴다

from server.app import report  # noqa: E402
from server.app.report import embedded_ids, split_body  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TITLE = "드럼세탁기 탈수 진동 저감을 위한 스프링·댐퍼 설계 최적화 종합 검토 보고서"
BODY = """# 드럼세탁기 탈수 진동 저감을 위한 스프링·댐퍼 설계 최적화 종합 검토 보고서

## 요약
- 진동이 **34%** 줄었습니다
- 검증이 더 필요합니다

## 결과
본문 문장입니다.
1. 첫째
2. 둘째

### 결과
같은 이름의 소제목입니다.

### **굵은** `코드` 제목
```python
## 코드 안의 제목
x = "<b>" & 1
![코드 안 삽입](artifact:999999)
```
| 항목 | 값 |
|---|---|
| 진동 | **3.1** |

## 결과
다시 같은 이름입니다.
3. 다른 목록
4. 다른 목록 둘째
"""


class Page(HTMLParser):
    """HTML에서 id·링크·자원 참조를 모은다."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.ids, self.links, self.resources, self.scripts = [], [], [], 0
        self.headings, self._h = [], None
        self.toc_links, self._in_nav = [], False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if a.get("id"):
            self.ids.append(a["id"])
        if tag == "script":
            self.scripts += 1
        if tag == "nav":
            self._in_nav = True
        if tag == "a" and a.get("href", "").startswith("#"):
            self.links.append(a["href"][1:])
            if self._in_nav:
                self.toc_links.append(a["href"][1:])
        if tag in ("img", "script", "iframe", "source", "video", "audio", "embed") and a.get("src"):
            self.resources.append(a["src"])
        if tag == "link" and a.get("href"):
            self.resources.append(a["href"])
        if tag in ("h2", "h3") and a.get("id"):
            self._h = [a["id"], ""]

    def handle_endtag(self, tag):
        if tag == "nav":
            self._in_nav = False
        if tag in ("h2", "h3") and self._h:
            self.headings.append(tuple(self._h))
            self._h = None

    def handle_data(self, data):
        if self._h:
            self._h[1] += data


def parse_page(html: str) -> Page:
    p = Page()
    p.feed(html)
    return p


def doc_parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def make_project(client, admin, name="진동 저감 과제"):  # noqa: F811
    return client.post("/api/projects", json={"name": name, "description": "탈수 진동 과제"}, headers=admin).json()["id"]


def add_results(client, admin, pid):  # noqa: F811
    tbl = client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={
        "kind": "table", "title": "케이스", "payload": {"columns": ["ID", "f"], "rows": [[1, 0.5], [2, 0.4]]}}).json()
    line = client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={
        "kind": "line", "title": "수렴", "payload": {"x": "회", "y": "f", "series": [{"name": "f", "points": [{"x": 1, "y": 2}, {"x": 2, "y": 1}]}]}}).json()
    return tbl["id"], line["id"]


@pytest.fixture(scope="module")
def admin(client):  # noqa: F811
    return _auth(client, "admin")[0]


@pytest.fixture(scope="module")
def ai(client, admin):  # noqa: F811
    """AI 보고서가 있는 프로젝트: (프로젝트 id, 보고서 응답, HTML, docx 바이트)"""
    pid = make_project(client, admin)
    tid, lid = add_results(client, admin, pid)
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "첫 판", "body": "## 첫 판"})
    body = BODY.replace("다시 같은 이름입니다.", f"다시 같은 이름입니다.\n![케이스 결과](artifact:{tid})\n![수렴 이력](artifact:{lid})")
    rep = client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": TITLE, "body": body}).json()
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    docx = client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content
    return SimpleNamespace(pid=pid, rep=rep, html=html, docx=docx)


@pytest.fixture(scope="module")
def fallback(client, admin):  # noqa: F811
    """AI 보고서가 없을 때 게시된 결과를 모은 것"""
    pid = make_project(client, admin, "모음 과제")
    client.post(f"/api/projects/{pid}/flows", headers=admin, json={
        "title": "흐름", "goal": "진동 저감", "nodes": [{"id": "a", "label": "해석", "type": "solve"}], "edges": []})
    add_results(client, admin, pid)
    client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={"kind": "markdown", "title": "결론", "payload": {"text": "## 소제목\n내용"}})
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    docx = client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content
    return SimpleNamespace(pid=pid, html=html, docx=docx)


# ---------- 틀 ----------
def _assert_html_frame(html: str, pid: int, doc_no: str):
    page = parse_page(html)
    assert re.fullmatch(r"DDS-P\d{4}-(R\d{3}|AUTO)", doc_no) and f"DDS-P{pid:04d}-" in doc_no
    for cls in ("doc-head", "cover", "colophon", "doc-foot"):
        assert f'class="{cls}"' in html, cls
    assert html.count(doc_no) >= 3  # 머리글·표지·꼬리말
    assert "DDS Conversa" in html and "담당 엔지니어" in html and "<title>" in html
    assert 'rel="icon"' in html and "data:image/svg+xml;base64," in html
    assert '<svg class="dds-logo"' in html
    assert page.scripts == 1  # 틀이 넣은 현재 위치 표시 하나뿐이다
    return page


def test_ai_report_html_frame(ai):
    page = _assert_html_frame(ai.html, ai.pid, ai.rep["doc_no"])
    assert ai.rep["doc_no"] == f"DDS-P{ai.pid:04d}-R{ai.rep['id']:03d}"
    assert ai.rep["edition"] == 2 and "2판" in ai.html  # 프로젝트 안에서 두 번째 판
    assert ai.rep["author"] and ai.rep["author"] in ai.html
    assert ai.rep["issued_at"] in ai.html and ai.rep["project_name"] in ai.html
    assert ai.rep["provenance"] in ai.html and ai.rep["colophon"] in ai.html
    assert page.toc_links


def test_fallback_html_frame(fallback):
    page = _assert_html_frame(fallback.html, fallback.pid, f"DDS-P{fallback.pid:04d}-AUTO")
    assert "모음 과제" in fallback.html and "자동" in fallback.html
    # 작업 흐름·결과 제목이 목차가 되고, 그 아래 소제목(결과 안의 제목)은 목차에 넣지 않는다
    names = {i: t for i, t in page.headings}
    assert [names[i] for i in page.toc_links] == ["작업 흐름 — 흐름", "결론", "케이스", "수렴"]
    assert "소제목" in fallback.html and "<h3>소제목</h3>" in fallback.html


def test_favicon_decodes_to_project_icon(ai):
    data = re.search(r'rel="icon"[^>]*href="data:image/svg\+xml;base64,([^"]+)"', ai.html).group(1)
    assert base64.b64decode(data).decode() == (ROOT / "server/app/assets/favicon.svg").read_text(encoding="utf-8")


def test_web_response_carries_same_frame(client, admin, ai):  # noqa: F811
    latest = client.get(f"/api/projects/{ai.pid}/reports/latest", headers=admin).json()
    one = client.get(f"/api/reports/{ai.rep['id']}", headers=admin).json()
    for r in (latest, one, ai.rep):
        assert r["doc_no"] == ai.rep["doc_no"] and r["edition"] == 2 and r["project_name"] == "진동 저감 과제"
        assert r["provenance"].startswith("DDS Conversa 발행") and "담당 엔지니어" in r["provenance"]
        # 제목 줄을 표지가 이미 보여 주므로 본문에서 뺀다
        assert not r["body"].lstrip().startswith("# ")


# ---------- 목차 ----------
def test_contents_anchors_resolve_and_are_unique(ai):
    page = parse_page(ai.html)
    assert len(page.ids) == len(set(page.ids))
    assert page.toc_links and set(page.toc_links) <= set(page.ids)
    names = {i: t for i, t in page.headings}
    # 같은 이름 제목 셋(결과·결과·결과)이 서로 다른 곳으로 간다
    assert [names[i] for i in page.toc_links] == ["요약", "결과", "결과", "굵은 코드 제목", "결과"]
    assert len(set(page.toc_links)) == 5


def test_contents_list_has_plain_text_and_levels(ai):
    nav = re.search(r"<nav.*?</nav>", ai.html, re.S).group(0)
    assert "<strong>" not in nav and "<code>" not in nav and "**" not in nav and "`" not in nav
    assert 'class="l3"' in nav and 'class="l2"' in nav


def test_code_fence_is_not_a_heading(ai):
    assert "코드 안의 제목" not in re.search(r"<nav.*?</nav>", ai.html, re.S).group(0)
    assert "## 코드 안의 제목" in ai.html  # 코드 그대로 보인다
    assert 'x = &quot;&lt;b&gt;&quot; &amp; 1' in ai.html or 'x = "&lt;b&gt;" &amp; 1' in ai.html
    assert "코드 안 삽입" in ai.html


def test_embed_inside_code_fence_is_not_a_result():
    body = "## 가\n```\n![예](artifact:1)\n```\n![진짜](artifact:2)\n"
    assert embedded_ids(body) == [2]
    assert [p[2] for p in split_body(body) if p[0] == "art"] == [2]


def test_title_not_printed_twice(ai):
    page_body = ai.html.split('class="cover"', 1)[1]
    assert page_body.count(TITLE) == 1


def test_other_top_level_heading_is_a_section():
    d = report.compose(SimpleNamespace(id=1, name="p", description=""), [], None,
                       SimpleNamespace(id=1, title="제목", body="# 제목\n# 다른 머리\n내용", created_at=datetime(2026, 1, 1)))
    assert [(t.level, t.text) for t in d.toc] == [(2, "다른 머리")]


def test_no_headings_is_still_valid(client, admin):  # noqa: F811
    pid = make_project(client, admin, "제목 없는 과제")
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "제목 없음", "body": "그냥 문단만 있습니다."})
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    assert "<nav" not in html and 'class="cover"' in html and "그냥 문단만 있습니다." in html
    data = client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content
    doc = Document(io.BytesIO(data))
    assert any("그냥 문단만" in p.text for p in doc.paragraphs) and not any(p.text == "목차" for p in doc.paragraphs)
    assert not doc_parts(data)["word/document.xml"].count(b"w:hyperlink")


def test_title_only_body_is_valid(client, admin):  # noqa: F811
    pid = make_project(client, admin, "제목만")
    rep = client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "짧은 제목", "body": "# 짧은 제목"}).json()
    assert rep["body"] == ""
    assert '<h1>짧은 제목</h1>' in client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    assert client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content[:2] == b"PK"


# ---------- 파일 하나로 끝난다 ----------
def _no_external(html: str):
    page = parse_page(html)
    assert not [r for r in page.resources if not r.startswith("data:")]
    without_links = re.sub(r"""<a\b[^>]*>""", "<a>", html)  # 본문의 링크(누르면 여는 것)는 자원이 아니다
    assert not re.search(r"https?://", without_links), re.search(r".{40}https?://.{40}", without_links)
    assert not re.search(r"""url\(\s*(?!['"]?data:)""", html) and "@import" not in html


def test_html_is_self_contained(ai, fallback):
    _no_external(ai.html)
    _no_external(fallback.html)


def test_external_image_in_body_is_not_loaded(client, admin):  # noqa: F811
    pid = make_project(client, admin, "외부 그림")
    body = "## 가\n![외부 그림](http://evil.example/a.png)\n[사내 링크](https://intranet.example/doc)"
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "t", "body": body})
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    _no_external(html)
    assert "evil.example" not in html and "[이미지: 외부 그림]" in html
    assert 'href="https://intranet.example/doc"' in html
    doc = client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content
    assert b"evil.example" not in doc_parts(doc)["word/_rels/document.xml.rels"]
    assert b"https://intranet.example/doc" in doc_parts(doc)["word/_rels/document.xml.rels"]


# ---------- 글 넣기 공격 ----------
EVIL = '</style><script>alert(1)</script><img src=x onerror=alert(2)>"\''


def test_ai_text_cannot_inject(client, admin):  # noqa: F811
    pid = client.post("/api/projects", json={"name": "프로젝트" + EVIL}, headers=admin).json()["id"]
    art = client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={
        "kind": "table", "title": "표" + EVIL, "payload": {"columns": [EVIL], "rows": [[EVIL]]}}).json()["id"]
    body = (f"## 제목 {EVIL}\n<script>alert(3)</script>\n\n<div onclick=alert(4)>글</div>\n\n"
            f"[누르면](javascript:alert(5))\n\n[부호화](&#106;avascript:alert(6))\n\n"
            f"![캡션 <script>alert(7)</script>](artifact:{art})\n\n```\n<script>alert(8)</script>\n```\n")
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "제목" + EVIL, "body": body})
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    page = parse_page(html)
    assert page.scripts == 1 and "alert" not in html.split("<script>")[1].split("</script>")[0]
    without_own = html.replace(re.search(r"<script>.*?</script>", html, re.S).group(0), "")
    assert "<script" not in without_own
    assert not re.search(r"<(img|div|a)\b[^>]*\bon\w+=", without_own)
    hrefs = re.findall(r"""href=["']([^"']*)""", html)
    assert not [h for h in hrefs if "javascript" in h.lower() or "&#106;" in h]
    assert html.count("</style>") == 1
    # 막힌 글은 사라지지 않고 글자로 보인다
    assert "&lt;script&gt;alert(3)&lt;/script&gt;" in html and "&lt;script&gt;alert(8)&lt;/script&gt;" in html
    data = client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content
    doc = Document(io.BytesIO(data))
    assert any("<script>alert(3)</script>" in p.text for p in doc.paragraphs)
    assert b"javascript" not in doc_parts(data)["word/_rels/document.xml.rels"]
    # 목차 글자도 글자로만 들어간다
    root = etree.fromstring(doc_parts(data)["word/document.xml"])
    assert any("<script>" in "".join(h.itertext()) for h in root.iter(f"{W}hyperlink"))
    assert client.get(f"/api/projects/{pid}/reports/latest", headers=admin).status_code == 200


def test_flow_status_cannot_inject():
    flow = SimpleNamespace(title="흐름", goal="", nodes=[{"id": "a", "label": "가", "status": "x'><script>alert(1)</script>"}], edges=[])
    html = report.build_html(SimpleNamespace(id=1, name="p", description=""), [], flow, None)
    assert parse_page(html).scripts == 1


def test_control_characters_do_not_break_word():
    data = report.build_docx(SimpleNamespace(id=1, name="p\x01", description=""), [], None,
                             SimpleNamespace(id=1, title="제목\x01\x0b", body="## 가\x01나\n본문\x02", created_at=datetime(2026, 1, 1)))
    assert Document(io.BytesIO(data)).core_properties.title.startswith("제목")


# ---------- Word ----------
def _xml(data: bytes, name="word/document.xml"):
    return etree.fromstring(doc_parts(data)[name])


def test_word_bookmarks_and_hyperlinks_match(ai, fallback):
    for data in (ai.docx, fallback.docx):
        root = _xml(data)
        marks = [b.get(f"{W}name") for b in root.iter(f"{W}bookmarkStart")]
        ids = [b.get(f"{W}id") for b in root.iter(f"{W}bookmarkStart")]
        ends = [b.get(f"{W}id") for b in root.iter(f"{W}bookmarkEnd")]
        anchors = [h.get(f"{W}anchor") for h in root.iter(f"{W}hyperlink") if h.get(f"{W}anchor")]
        assert anchors and anchors == marks and len(set(marks)) == len(marks)
        assert sorted(ids) == sorted(ends) and len(set(ids)) == len(ids)
        for name in marks:  # Word 책갈피 이름 규칙: 글자로 시작, 40자 이내, 글자·숫자·밑줄
            assert re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", name)


def test_word_contents_text_matches_headings(ai):
    root = _xml(ai.docx)
    toc = {h.get(f"{W}anchor"): "".join(h.itertext()) for h in root.iter(f"{W}hyperlink") if h.get(f"{W}anchor")}
    assert list(toc.values()) == ["요약", "결과", "결과", "굵은 코드 제목", "결과"]
    # 책갈피가 그 제목 문단 안에 있다
    for bm in root.iter(f"{W}bookmarkStart"):
        par = bm.getparent()
        assert par.find(f"{W}pPr/{W}pStyle").get(f"{W}val").startswith("Heading")
        assert "".join(par.itertext()) == toc[bm.get(f"{W}name")]
    assert list(toc.keys()) == parse_page(ai.html).toc_links  # HTML과 같은 이름


def test_word_asks_no_field_update(ai):
    assert b"updateFields" not in doc_parts(ai.docx)["word/settings.xml"]


def test_word_opens_without_compatibility_mode(ai):
    root = _xml(ai.docx, "word/settings.xml")
    modes = [s.get(f"{W}val") for s in root.iter(f"{W}compatSetting") if s.get(f"{W}name") == "compatibilityMode"]
    assert modes == ["15"]


def test_word_paragraph_properties_are_in_schema_order(ai):
    """순서가 어긋나면 Word가 파일이 손상됐다고 묻는다."""
    for name in ("word/document.xml", "word/styles.xml", *[k for k in doc_parts(ai.docx) if re.fullmatch(r"word/(header|footer)\d*\.xml", k)]):
        for ppr in _xml(ai.docx, name).iter(f"{W}pPr"):
            order = [report._PPR_ORDER.index("w:" + ch.tag.split("}")[1]) for ch in ppr]
            assert order == sorted(order), [ch.tag for ch in ppr]


def _section_xml(data, kind):
    parts = doc_parts(data)
    return b"".join(v for k, v in parts.items() if re.fullmatch(rf"word/{kind}\d*\.xml", k))


def test_word_header_footer(ai, fallback):
    for data, doc_no in ((ai.docx, ai.rep["doc_no"]), (fallback.docx, f"DDS-P{fallback.pid:04d}-AUTO")):
        head, foot = _section_xml(data, "header"), _section_xml(data, "footer")
        assert "DDS Conversa".encode() in head and b"pic:pic" in head
        assert doc_no.encode() in foot and b" PAGE " in foot and b" NUMPAGES " in foot
        assert Document(io.BytesIO(data)).sections[0].different_first_page_header_footer  # 표지는 머리글 없이 큰 로고로


def test_word_logo_is_project_png(ai):
    parts = doc_parts(ai.docx)
    logo = (ROOT / "server/app/assets/dds-logo.png").read_bytes()
    assert any(k.startswith("word/media/") and v == logo for k, v in parts.items())


def test_word_core_properties(ai, fallback):
    props = Document(io.BytesIO(ai.docx)).core_properties
    assert props.author == "DDS Conversa" and props.title == TITLE and "python-docx" not in (props.comments or "")
    assert props.last_modified_by == "DDS Conversa" and ai.rep["doc_no"] in props.subject
    props = Document(io.BytesIO(fallback.docx)).core_properties
    assert props.author == "DDS Conversa" and props.title == "모음 과제"


def test_word_table_style_exists(ai, fallback):
    for data in (ai.docx, fallback.docx):
        doc = Document(io.BytesIO(data))
        names = {s.style_id for s in doc.styles}
        used = {t.get(f"{W}val") for t in _xml(data).iter(f"{W}tblStyle")}
        assert used and used <= names


def test_word_title_and_cover(ai):
    doc = Document(io.BytesIO(ai.docx))
    texts = [p.text for p in doc.paragraphs]
    assert texts.count(TITLE) == 1
    assert "목차" in texts
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    for v in (ai.rep["doc_no"], "2판", ai.rep["issued_at"], ai.rep["author"]):
        assert any(v in c for c in cells), v
    assert any(ai.rep["provenance"] in t for t in texts) and any(ai.rep["colophon"] in t for t in texts)


def test_word_code_fence_and_inline(ai):
    doc = Document(io.BytesIO(ai.docx))
    heads = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert not any("코드 안의 제목" in p.text for p in heads)
    assert any(p.text.startswith("## 코드 안의 제목") and 'x = "<b>" & 1' in p.text for p in doc.paragraphs)
    assert any(r.bold and r.text == "34%" for p in doc.paragraphs for r in p.runs)
    assert not any("**" in p.text for p in doc.paragraphs)


def test_word_numbered_lists_restart(ai):
    doc = Document(io.BytesIO(ai.docx))
    nums = [p._p.pPr.numPr.numId.val for p in doc.paragraphs
            if p.style.name == "List Number" and p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(nums) == 4 and nums[0] == nums[1] and nums[2] == nums[3] and nums[0] != nums[2]


# ---------- 원래 동작 유지 ----------
def test_figure_numbering_and_deleted_result(client, admin):  # noqa: F811
    pid = make_project(client, admin, "번호")
    tid, lid = add_results(client, admin, pid)
    gone = client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={
        "kind": "line", "title": "지울 것", "payload": {"series": [{"name": "f", "points": [{"x": 1, "y": 1}]}]}}).json()["id"]
    body = f"## 결과\n![지운 그림](artifact:{gone})\n![수렴 이력](artifact:{lid})\n![케이스 결과](artifact:{tid})"
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "t", "body": body})
    assert client.delete(f"/api/artifacts/{gone}", headers=admin).status_code == 200
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    assert "그림 1. 수렴 이력" in html and "표 1. 케이스 결과" in html and "지운 그림" not in html and "그림 2." not in html
    doc = Document(io.BytesIO(client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content))
    texts = [p.text for p in doc.paragraphs]
    assert "그림 1. 수렴 이력" in texts and "표 1. 케이스 결과" in texts and "지운 그림" not in " ".join(texts)


def test_table_row_limit_kept(client, admin):  # noqa: F811
    pid = make_project(client, admin, "큰 표")
    tid = client.post(f"/api/projects/{pid}/artifacts", headers=admin, json={
        "kind": "table", "title": "큰 표", "payload": {"columns": ["n"], "rows": [[i] for i in range(250)]}}).json()["id"]
    client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "t", "body": f"![큰 표](artifact:{tid})"})
    html = client.get(f"/api/projects/{pid}/report?format=html", headers=admin).text
    assert ">199<" in html and ">200<" not in html
    doc = Document(io.BytesIO(client.get(f"/api/projects/{pid}/report?format=docx", headers=admin).content))
    assert len(doc.tables[-1].rows) == 201  # 머리 + 200행


# ---------- 틀과 앱이 어긋나지 않게 ----------
def _light_tokens() -> dict[str, str]:
    css = (ROOT / "frontend/src/design/tokens.css").read_text(encoding="utf-8")
    light = re.split(r'\n:root\[data-theme="dark"\] \{', css, maxsplit=1)[0]
    return {k: v.lower() for k, v in re.findall(r"(--[a-z0-9-]+)\s*:\s*(#[0-9a-fA-F]{6})\b", light)}


def test_report_colors_match_tokens():
    tokens = _light_tokens()
    for name, value in report.COLOR.items():
        assert tokens[f"--{name}"] == value.lower(), name
    # 그래프 계열은 앱의 chartColors()와 같은 순서(주홍 → 청록 → 보라 → 호박 → 회청 → 빨강)
    order = ["--brand", "--mean", "--sigma", "--warn-solid", "--epi", "--err"]
    assert [p.lower() for p in report.PALETTE] == [tokens[t] for t in order]
    plot = (ROOT / "frontend/src/components/Plot.tsx").read_text(encoding="utf-8")
    assert '["--brand", "--mean", "--sigma", "--warn-solid", "--epi", "--err"]' in plot
    assert "#2563eb" not in (ROOT / "server/app/report.py").read_text(encoding="utf-8")


def test_css_uses_only_token_colors(ai):
    tokens = set(_light_tokens().values())
    css = re.search(r"<style>(.*?)</style>", ai.html, re.S).group(1)
    assert {c.lower() for c in re.findall(r"#[0-9a-fA-F]{6}\b", css)} <= tokens


def test_server_assets_equal_frontend_public():
    """Linux 서버에 frontend/public이 없어도 되게 서버 쪽에 사본을 두고, 어긋나지 않게 막는다."""
    server, front = ROOT / "server/app/assets", ROOT / "frontend/public"
    pairs = {"favicon.svg": front / "favicon.svg", "dds-logo.svg": front / "brand/dds-logo.svg", "dds-logo.png": front / "brand/dds-logo.png"}
    for name, source in pairs.items():
        a, b = (server / name).read_bytes(), source.read_bytes()
        assert a.replace(b"\r\n", b"\n") == b.replace(b"\r\n", b"\n"), name


def test_server_files_are_in_git():
    """서버는 git pull로 배포한다. report_md.py가 빠지면 서버가 아예 뜨지 못하고, 그림 사본이 빠지면 모든 보고서가 500이다."""
    git = shutil.which("git")
    if not git or not (ROOT / ".git").exists():
        pytest.skip("git 저장소가 아닙니다")
    need = ["server/app/report_md.py", "server/app/assets/favicon.svg", "server/app/assets/dds-logo.svg", "server/app/assets/dds-logo.png"]
    listed = subprocess.run([git, "ls-files", "--", *need], cwd=ROOT, capture_output=True, text=True).stdout.split()
    assert sorted(listed) == sorted(need), "git add가 필요합니다"


def test_instructions_say_body_only():
    ins = (ROOT / "agent-context/INSTRUCTIONS.md").read_text(encoding="utf-8")
    section = ins.split("## 1-0.", 1)[1].split("\n## ", 1)[0]
    for word in ("본문만", "표지", "목차", "로고", "`##`"):
        assert word in section, word
    tool = (ROOT / "mcp_servers/dds_flow.py").read_text(encoding="utf-8")
    doc = tool.split("def write_report", 1)[1].split('"""')[1]
    for word in ("본문만", "표지", "목차", "로고", "##"):
        assert word in doc, word
    spec = (ROOT / "docs/product-spec.md").read_text(encoding="utf-8")
    assert "Word·HTML 파일의 모양은 바꾸지 않는다" not in spec and "문서 번호" in spec and "server/app/assets" in spec
    assert "server/app/assets" in (ROOT / "AGENTS.md").read_text(encoding="utf-8")


def test_demo_agent_writes_body_only():
    src = (ROOT / "agent/dds_agent/demo_agent.py").read_text(encoding="utf-8")
    assert '"## 요약\\n"' in src and 'lines = ["## 요약"]' in src
    assert not re.search(r'["\']# ', src)


# ---------- 파일로 나갈 때 값이 사라지거나 달라지면 안 된다 ----------
def _project(pid=1, name="p", description=""):
    return SimpleNamespace(id=pid, name=name, description=description)


def _report(body, title="제목"):
    return SimpleNamespace(id=1, title=title, body=body, created_at=datetime(2026, 1, 1))


def _art(aid, kind, payload, title="결과"):
    return SimpleNamespace(id=aid, kind=kind, title=title, payload=payload, created_at=datetime(2026, 1, 1))


def _doe(n_cols, n_rows=3, name="param_{:02d}_len"):
    return {"columns": [name.format(i) for i in range(n_cols)], "rows": [[11.111 + r + i for i in range(n_cols)] for r in range(n_rows)]}


def _sections(data):
    return list(_xml(data).iter(f"{W}sectPr"))


def _orient(sect):
    return sect.find(f"{W}pgSz").get(f"{W}orient") or "portrait"  # 세로는 생략된다


def _word_tables(data):
    """표지 표를 뺀 [(칸 글자 2차원 목록, 열 폭 합(twip))]"""
    out = []
    for tbl in _xml(data).iter(f"{W}tbl"):
        rows = [["".join(tc.itertext()) for tc in tr.findall(f"{W}tc")] for tr in tbl.findall(f"{W}tr")]
        if rows[0][0] != "문서 번호":
            out.append((rows, sum(int(g.get(f"{W}w")) for g in tbl.find(f"{W}tblGrid"))))
    return out


def test_wide_table_prints_landscape_with_repeating_head():
    html = report.build_html(_project(), [_art(5, "table", _doe(12))], None, _report("## 넓은 표\n![넓은](artifact:5)\n\n## 다음\n글"))
    main = html.split("<main>")[1]
    assert '<figure class="wide">' in main and '<div class="tw wide' in main
    assert "<thead>" in main and "@page wide{size:A4 landscape}" in html
    assert "page:wide" in html.split("@media print")[1] and "has(+figure.wide)" in html
    assert "overflow-wrap:anywhere" in html.split("pre{white-space:pre-wrap")[1].split("}")[0]


def test_narrow_table_stays_portrait():
    html = report.build_html(_project(), [_art(5, "table", {"columns": ["ID", "f"], "rows": [[1, 0.5]]})], None, _report("![t](artifact:5)"))
    assert "<figure><div class=\"tw\"><table><thead>" in html


def test_too_wide_table_is_split_for_print_with_all_values():
    html = report.build_html(_project(), [_art(5, "table", _doe(40))], None, _report("![넓은](artifact:5)"))
    split = html.split('<div class="print-split">')[1].split("<figcaption>")[0]
    assert split.count("<table") >= 2 and "첫 열은 반복" in split
    for i in range(40):
        assert f"param_{i:02d}_len" in split, i
    assert split.count(">param_00_len<") == split.count("<table")  # 첫 열이 표마다 반복된다
    assert 'class="tw wide f7 screen-part"' in html  # 화면은 표 하나를 가로로 스크롤한다


def test_word_wide_table_goes_landscape_and_back():
    a = _art(5, "table", _doe(12))
    data = report.build_docx(_project(), [a], None, _report("## 앞\n글\n\n## 넓은 표\n![넓은](artifact:5)\n\n## 뒤\n글"))
    sects = _sections(data)
    assert [_orient(x) for x in sects] == ["portrait", "landscape", "portrait"]
    assert all(x.find(f"{W}titlePg") is None or x.find(f"{W}titlePg").get(f"{W}val") in ("0", "false") for x in sects[1:])
    doc = Document(io.BytesIO(data))
    # 표 바로 앞 제목이 앞 쪽에 혼자 남지 않고 표와 같은 절로 간다
    assert abs(doc.sections[1].header.paragraphs[0].paragraph_format.tab_stops[0].position - report.Mm(247)) < 1000
    body = list(_xml(data).find(f"{W}body"))
    heading = [i for i, el in enumerate(body) if "".join(el.itertext()) == "넓은 표"][-1]  # 앞의 하나는 목차 줄이다
    mark = next(i for i, el in enumerate(body) if el.find(f"{W}pPr/{W}sectPr") is not None)
    assert mark < heading and body[heading + 1].tag == f"{W}tbl"
    head, foot = _section_xml(data, "header"), _section_xml(data, "footer")
    assert b"pic:pic" in head and b" NUMPAGES " in foot
    # 예전 값 그대로 다 있다
    rows, width = _word_tables(data)[-1]
    assert len(rows[0]) == 12 and rows[1][0] == "11.111"
    assert width <= report.Mm(247).twips + 2


def test_word_consecutive_wide_tables_have_no_blank_page():
    """넓은 표 둘 사이에 제목만 있으면 세로 쪽으로 돌아가지 않는다. 돌아가면 빈 쪽이 생긴다."""
    arts = [_art(5, "table", _doe(12)), _art(6, "table", _doe(12))]
    data = report.build_docx(_project(), arts, None, _report("## 가\n![첫째](artifact:5)\n\n### 나\n![둘째](artifact:6)\n"))
    assert [_orient(x) for x in _sections(data)] == ["portrait", "landscape", "portrait"]


def test_word_heading_follows_content_back_to_portrait():
    data = report.build_docx(_project(), [_art(5, "table", _doe(12))], None, _report("![넓은](artifact:5)\n\n## 뒤\n글\n"))
    body = list(_xml(data).find(f"{W}body"))
    marks = [i for i, el in enumerate(body) if el.find(f"{W}pPr/{W}sectPr") is not None]
    heading = [i for i, el in enumerate(body) if "".join(el.itertext()) == "뒤"][-1]
    assert marks[-1] < heading  # 제목 '뒤'는 가로 쪽 끝에 남지 않고 글과 함께 세로 쪽으로 간다


def test_word_columns_hold_numbers_on_one_line():
    """좁은 칸에서 값이 '11.' '111'로 갈라지면 잘못 읽는다: 칸 폭이 값의 폭보다 넓어야 한다."""
    for n, cols in ((3, 3), (12, 12), (18, None), (70, None)):
        table = _doe(n)
        plan = report.plan_table([[str(c) for c in table["columns"]], *[[report._fmt(v) for v in r] for r in table["rows"]]], 160, 247)
        for group in plan.groups:
            widths = plan.widths(group, 247 if plan.landscape else 160)
            assert sum(widths) <= (247 if plan.landscape else 160) + 0.01
            for c, w in zip(group, widths):
                # 머리글 'param_00_len'(12자)이 한 줄에 들어갈 폭
                assert w >= 12 * 0.55 * plan.font * 0.3528, (n, c, w)
        if cols:
            assert len(plan.groups) == 1 and len(plan.groups[0]) == cols


def test_word_too_wide_table_repeats_first_column():
    data = report.build_docx(_project(), [_art(5, "table", _doe(40))], None, _report("![넓은](artifact:5)"))
    parts = _word_tables(data)
    assert len(parts) >= 2 and all(rows[0][0] == "param_00_len" for rows, _ in parts)
    seen = {name for rows, _ in parts for name in rows[0]}
    assert seen == {f"param_{i:02d}_len" for i in range(40)}
    assert any("첫 열은 반복" in p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_margin_logo_on_printed_pages(ai):
    uri = re.search(r'@top-left\{content:url\("data:image/svg\+xml;base64,([^"]+)"\)', ai.html).group(1)
    svg = base64.b64decode(uri).decode()
    assert 'width="11"' in svg and "xmlns" in svg and "#1428A0" in svg


def test_toc_script_follows_current_section():
    assert "a.offsetTop-box.offsetTop" not in report._SCRIPT and "var t=a.offsetTop;" in report._SCRIPT
    assert "'pointerdown'" in report._SCRIPT


class TestPictures:
    def _png(self, w, h):
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (w, h), "white").save(buf, "PNG")
        return buf.getvalue()

    def _extent(self, monkeypatch, w, h):
        monkeypatch.setattr(report, "image_bytes", lambda pid, a: self._png(w, h))
        data = report.build_docx(_project(), [_art(5, "image", {"path": "a.png"})], None, _report("![그림](artifact:5)"))
        # 표지 로고 · 이 그림 · 맺음말 로고 순서다
        ext = list(_xml(data).iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"))[1]
        return int(ext.get("cx")) / 36000, int(ext.get("cy")) / 36000  # mm

    def test_tall_image_fits_page(self, monkeypatch):
        w, h = self._extent(monkeypatch, 800, 3000)
        assert h <= report.IMAGE_MAX_MM + 0.1 and abs(w / h - 800 / 3000) < 0.01

    def test_wide_image_fits_width(self, monkeypatch):
        w, h = self._extent(monkeypatch, 3000, 600)
        assert w <= report.PORTRAIT_MM + 0.1 and abs(w / h - 5) < 0.01

    def test_small_image_is_not_stretched(self, monkeypatch):
        w, _h = self._extent(monkeypatch, 160, 120)
        assert abs(w - 160 / 96 * 25.4) < 0.1

    def test_print_limits_image_height(self, ai):
        assert "max-height:200mm" in ai.html

    def test_missing_image_says_so(self, monkeypatch):
        monkeypatch.setattr(report, "image_bytes", lambda pid, a: None)
        rep = _report("![사라진 그림](artifact:5)")
        arts = [_art(5, "image", {"path": "gone.png"})]
        html = report.build_html(_project(), arts, None, rep)
        assert report.NO_IMAGE in html and "그림 1. 사라진 그림" in html
        texts = [p.text for p in Document(io.BytesIO(report.build_docx(_project(), arts, None, rep))).paragraphs]
        at = texts.index(report.NO_IMAGE)
        assert texts[at + 1] == "그림 1. 사라진 그림"

    def test_unreadable_image_does_not_stop_word(self, monkeypatch):
        monkeypatch.setattr(report, "image_bytes", lambda pid, a: b"not an image")
        data = report.build_docx(_project(), [_art(5, "image", {"path": "a.png"})], None, _report("![그림](artifact:5)"))
        assert any(p.text == report.NO_WORD_IMAGE for p in Document(io.BytesIO(data)).paragraphs)


class TestCharts:
    @pytest.mark.parametrize("label", ["변형률(%) ($$)", "$\\\\응력$", "a$b_$c", "$\\\\frob{$"])
    def test_dollar_labels_stay_text(self, label):
        a = _art(8, "line", {"x": label, "y": label, "series": [
            {"name": label, "points": [{"x": 1, "y": 2}]}, {"name": "둘", "points": [{"x": 1, "y": 1}]}]})
        assert report.chart_png(a)[:4] == b"\x89PNG"
        b = _art(9, "bar", {"categories": [label, "b"], "series": [{"name": label, "values": [1, 2]}]})
        assert report.chart_png(b)[:4] == b"\x89PNG"

    def test_broken_chart_does_not_stop_report(self):
        arts = [_art(8, "line", {"series": [{"name": "x", "points": [{"nope": 1}]}]}), _art(9, "bar", {"categories": ["a", "b"], "series": [{"values": [1, 2, 3]}]})]
        rep = _report("## 가\n![망가진 선](artifact:8)\n![망가진 막대](artifact:9)\n뒤의 글")
        html = report.build_html(_project(), arts, None, rep)
        assert html.count(report.NO_CHART) == 2 and "뒤의 글" in html and "DDS Conversa" in html
        assert sum(p.text == report.NO_CHART for p in Document(io.BytesIO(report.build_docx(_project(), arts, None, rep))).paragraphs) == 2

    @pytest.mark.parametrize("a", [_art(8, "line", {"series": []}), _art(8, "scatter", {"series": [{"points": []}]}), _art(8, "bar", {"categories": [], "series": []})])
    def test_empty_chart_has_no_picture(self, a):
        assert report.chart_png(a) is None

    def _ticks(self, monkeypatch, cats):
        seen = {}
        real = report._set_xticks

        def spy(fig, ax, ticks, names):
            real(fig, ax, ticks, names)
            renderer = fig.canvas.get_renderer()
            labels = [t for t in ax.get_xticklabels() if t.get_text()]
            seen["labels"] = labels
            seen["rotation"] = labels[0].get_rotation()
            seen["boxes"] = [t.get_window_extent(renderer) for t in labels]
        monkeypatch.setattr(report, "_set_xticks", spy)
        report.chart_png(_art(8, "bar", {"categories": cats, "series": [{"name": "s", "values": [1] * len(cats)}]}))
        return seen

    def test_few_labels_stay_flat(self, monkeypatch):
        assert self._ticks(monkeypatch, ["리브", "플랜지", "보스"])["rotation"] == 0

    def test_many_long_labels_tilt(self, monkeypatch):
        seen = self._ticks(monkeypatch, [f"spring_position_{i}" for i in range(14)])
        assert seen["rotation"] == 45 and len(seen["labels"]) == 14

    def test_very_many_labels_thin_out(self, monkeypatch):
        seen = self._ticks(monkeypatch, [f"c{i:03d}_long_name" for i in range(90)])
        assert seen["rotation"] == 90 and len(seen["labels"]) < 90
        xs = sorted(b.x0 for b in seen["boxes"])
        assert all(q - p >= 0 for p, q in zip(xs, xs[1:]))


class TestCover:
    def test_long_title_uses_smaller_type(self):
        title = "긴 제목 " * 40
        data = report.build_docx(_project(), [], None, _report("## 가", title.strip()[:200]))
        doc = Document(io.BytesIO(data))
        par = next(p for p in doc.paragraphs if p.style.name == "Title")
        assert par.runs[0].font.size.pt == 15
        assert 'class="cover xlong"' in report.build_html(_project(), [], None, _report("## 가", title.strip()[:200]))

    def test_short_title_keeps_size(self):
        doc = Document(io.BytesIO(report.build_docx(_project(), [], None, _report("## 가", "짧은 제목"))))
        assert next(p for p in doc.paragraphs if p.style.name == "Title").runs[0].font.size.pt == 26

    def test_long_project_name_cut_on_cover(self):
        d = report.compose(_project(name="과제" * 200), [], None, _report("## 가"))
        assert len(d.project) == report.COVER_PROJECT and d.project.endswith("…")

    def test_long_description_moves_to_body(self):
        desc = "설명 " * 300
        d = report.compose(_project(description=desc), [], None, None)
        assert len(d.summary) <= report.COVER_SUMMARY and d.summary.endswith("…")
        assert d.blocks[0].kind == "note" and d.blocks[0].text == desc  # 잘린 글은 본문에 온전히 남는다

    def test_short_description_stays_on_cover(self):
        d = report.compose(_project(description="짧은 설명"), [], None, None)
        assert d.summary == "짧은 설명" and not any(b.text == "짧은 설명" for b in d.blocks)


class TestBodyMatchesWebTab:
    """같은 글은 웹 화면(marked)과 파일에서 같은 구조여야 한다."""

    def _toc(self, body):
        toc = report.report_md.Toc()
        for kind, text, _ in split_body(report.clean_body(body, "제목")):
            if kind == "md":
                report.report_md.parse(text, toc)
        return [t.text for t in toc.items]

    def test_hash_without_space_is_text(self):
        body = "다음 결과입니다.\n#3 케이스가 가장 좋았습니다.\n\n#12 케이스는 실패했습니다.\n##요약\n####### 일곱\n\n## 진짜"
        assert self._toc(body) == ["진짜"]
        html = report.build_html(_project(), [], None, _report(body))
        assert "#3 케이스가 가장 좋았습니다." in html and "#12 케이스는 실패했습니다." in html
        texts = [p.text for p in Document(io.BytesIO(report.build_docx(_project(), [], None, _report(body)))).paragraphs]
        assert any("#3 케이스" in t for t in texts) and "진짜" in texts

    def test_indented_heading_is_heading(self):
        assert self._toc("글\n  ## 들여쓴 제목\n내용") == ["들여쓴 제목"]

    def test_heading_after_pipeless_table(self):
        body = "a | b\n---|---\n1 | 2\n## 뒤"
        assert self._toc(body) == ["뒤"]
        html = report.build_html(_project(), [], None, _report(body))
        assert "<nav" in html and "<td class=\"n\">1</td>" in html and "<td>## 뒤" not in html

    def test_pipeless_table_after_paragraph(self):
        html = report.report_md.parse("문단\na | b\n---|---\n1 | 2").html
        assert "<p>문단</p>" in html and "<table>" in html

    def test_ordered_list_keeps_its_numbers(self):
        html = report.report_md.parse("## 앞\n3. 셋째\n4. 넷째\n").html
        assert '<ol start="3">' in html
        assert "<ol>" in report.report_md.parse("1. 첫째\n2. 둘째").html
        data = report.build_docx(_project(), [], None, _report("## 앞\n3. 셋째\n4. 넷째\n"))
        numbering = doc_parts(data)["word/numbering.xml"]
        assert re.search(rb'w:startOverride w:val="3"', numbering)

    def test_mixed_list_types_do_not_merge(self):
        html = report.report_md.parse("- 가\n- 나\n1. 다\n2. 라").html
        assert "<ul>" in html and "<ol>" in html and "1. 다" not in html

    def test_crlf_fence_and_embed(self):
        body = "## 앞\r\n```\r\n## 코드\r\n![x](artifact:1)\r\n## 코드둘\r\n```\r\n## 뒤\r\n"
        assert embedded_ids(body) == []
        assert self._toc(body) == ["앞", "뒤"]

    def test_line_separator_is_not_a_line_break(self):
        assert self._toc("문단\u2028## 가짜") == []
        assert report.clean_body("# 제목\u2028글", "제목") == "# 제목\u2028글"

    def test_web_tab_uses_same_line_rule(self):
        src = (ROOT / "frontend/src/canvas/ReportView.tsx").read_text(encoding="utf-8")
        assert 'body.split(/\\r\\n|\\r|\\n/)' in src and 'body.split("\\n")' not in src


class TestLegacyBodies:
    def test_own_contents_section_is_removed(self):
        body = "# 자체 목차 보고서\n\n## 목차\n1. [요약](#요약)\n2. [결과](#결과)\n\n## 요약\n글\n\n## 결과\n글"
        cleaned = report.clean_body(body, "옛 판 자체 목차")
        d = report.compose(_project(), [], None, _report(body, "옛 판 자체 목차"))
        assert [t.text for t in d.toc] == ["요약", "결과"]
        assert "[요약](#요약)" not in cleaned and "**자체 목차 보고서**" in cleaned  # 다른 제목 글은 지우지 않는다

    def test_contents_section_with_prose_is_kept(self):
        body = "## 목차 구성 방법\n설명입니다.\n\n## 목차\n이 구역은 글입니다.\n\n## 뒤"
        assert report.clean_body(body, "t") == body

    def test_section_named_contents_with_text_is_kept(self):
        body = "## 목차\n이 보고서는 다음을 다룹니다.\n- 가\n\n## 뒤"
        assert report.clean_body(body, "t") == body

    def test_contents_like_code_is_kept(self):
        body = "```\n## 목차\n- 가\n```\n## 뒤"
        assert report.clean_body(body, "t") == body

    def test_web_response_has_no_own_contents(self, client, admin):  # noqa: F811
        pid = make_project(client, admin, "옛 판")
        body = "## 목차\n- 요약\n- 결과\n\n## 요약\n글"
        rep = client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "옛 판", "body": body}).json()
        assert rep["body"].startswith("## 요약")


class TestSafeText:
    BAD = "가\ufffe나\uffff다\x01"

    def test_all_text_paths(self):
        rep = _report(f"## {self.BAD}\n내용{self.BAD}\n\n| a{self.BAD} |\n|---|\n| b{self.BAD} |\n", f"제목{self.BAD}")
        flow = SimpleNamespace(title=self.BAD, goal=self.BAD, nodes=[{"id": "a", "label": self.BAD, "detail": self.BAD}], edges=[])
        art = _art(1, "table", {"columns": [self.BAD], "rows": [[self.BAD]]}, title=self.BAD)
        project = _project(name=f"프로젝트{self.BAD}", description=self.BAD)
        issue = report.make_issue(project, rep, 1, f"요청자{self.BAD}")
        data = report.build_docx(project, [art], flow, rep, issue)
        assert "\uffff" not in "".join(_xml(data).itertext())
        fallback = report.build_docx(project, [art], flow, None, report.make_issue(project, requester=self.BAD))
        assert fallback[:2] == b"PK"
        assert "\uffff" not in report.build_html(project, [art], flow, rep, issue)

    def test_route_returns_word(self, client, admin):  # noqa: F811
        pid = make_project(client, admin, "제어 문자")
        rid = client.post(f"/api/projects/{pid}/reports", headers=admin, json={"title": "제\uffff목", "body": "## 제\uffff목\n내용"}).json()["id"]
        r = client.get(f"/api/projects/{pid}/report?format=docx&report_id={rid}", headers=admin)
        assert r.status_code == 200 and r.content[:2] == b"PK"


def test_only_semicolon_entities_are_decoded():
    body = "## 주소 &copy=1\n\n서버 http://srv/api?a=1&copy=2&para=3 입니다. R&D, A&sect B. 기호 &copy; `x&amp;y`\n"
    data = report.build_docx(_project(), [], None, _report(body))
    text = " ".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "a=1&copy=2&para=3" in text and "A&sect B" in text and "기호 ©" in text and "x&amp;y" in text
    d = report.compose(_project(), [], None, _report(body))
    assert d.toc[0].text == "주소 &copy=1"
    assert "&amp;copy=2&amp;para=3" in report.build_html(_project(), [], None, _report(body))


def test_link_with_ampersand_keeps_address_in_word():
    data = report.build_docx(_project(), [], None, _report("[링크](http://srv/a?x=1&copy=2)"))
    assert b"copy=2" in doc_parts(data)["word/_rels/document.xml.rels"]


def test_web_markdown_escapes_quotes():
    src = (ROOT / "frontend/src/components/Markdown.tsx").read_text(encoding="utf-8")
    assert '.replace(/"/g, "&quot;")' in src and ".replace(/'/g, \"&#39;\")" in src


def test_large_word_table_is_not_quadratic():
    """row.cells는 부를 때마다 표 전체를 훑는다. 칸마다 부르면 넓은 표에서 Word 버튼이 수십 초 걸린다."""
    src = (ROOT / "server/app/report.py").read_text(encoding="utf-8")
    assert "row.cells[c]" not in src and "for tc, c, w in zip(tr.tc_lst" in src


def test_contents_limit_is_stated_in_spec():
    spec = (ROOT / "docs/product-spec.md").read_text(encoding="utf-8")
    assert "쪽 번호는 넣지 않는다" in spec and "가로 쪽" in spec
